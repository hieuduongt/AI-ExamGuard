"""
Đây là gì: script tạo file Word "Hướng dẫn xây dựng AI ExamGuard Python cho Codex"
  (build_huong_dan_codex.py).
Mục tiêu: sinh tự động bản nháp tài liệu hướng dẫn người không biết lập trình giao việc từng bước
  cho trợ lý AI Codex để tạo một ứng dụng Python chạy trên máy tính.
Chức năng: dựng tài liệu gồm trang bìa, mục 1 tổng quan, mục 2 năm câu lệnh mẫu, mục 3 bảng lệnh
  Terminal, mục 4 lời khuyên thuyết trình; tự định dạng phông chữ, bảng, danh sách, hộp ghi chú, đầu
  trang và số trang. Các câu lệnh mô tả kiến trúc cũ (main.py, thư mục ui/), không phải mã nguồn
  hiện tại. Thuật ngữ: python-docx là thư viện Python giúp tạo file Word (.docx); XML là "ngôn ngữ
  ghi chú" nằm bên trong file Word, các thẻ dạng "w:..." là từng mục cài đặt của Word; dxa là đơn vị
  đo của Word (1440 dxa = 1 inch, khoảng 2,54 cm); pt là đơn vị cỡ chữ.
Đầu ra: cùng một file Word được lưu ra hai nơi:
  docs/tai-lieu/Huong_Dan_Xay_Dung_AI_ExamGuard_Python_Cho_Codex.docx và docs/DU_AN_PYTHON/tai_lieu/
  (cùng tên file); thư mục gốc ROOT ở đây là docs/.
Dùng ở đâu: chạy tay bằng lệnh python3 docs/tools/build_huong_dan_codex.py (cần cài thư viện
  python-docx). Không chạy lại để đè lên bản trong docs/tai-lieu-du-thi/ vì bản đó đã được chỉnh
  tay.
Liên kết với: build_ho_so_du_an.py và build_mo_ta_tinh_nang.py (dùng chung bộ hàm định dạng);
  docs/van-de-ton-dong.md mục 8 và docs/README.md (giải thích vì sao không chạy lại).
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# Thư mục gốc (ở đây là docs/) và hai nơi lưu cùng một file Word.
ROOT = Path(__file__).resolve().parents[1]
OUTPUT_1 = ROOT / "tai-lieu" / "Huong_Dan_Xay_Dung_AI_ExamGuard_Python_Cho_Codex.docx"
OUTPUT_2 = ROOT / "DU_AN_PYTHON" / "tai_lieu" / "Huong_Dan_Xay_Dung_AI_ExamGuard_Python_Cho_Codex.docx"

# Bảng màu dùng chung cho tài liệu, viết bằng mã màu hex (6 ký tự, ví dụ "0B2545" là xanh navy).
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5C6B7A"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F2F7FC"
PALE_GREEN = "EDF8F2"
GREEN = "16794A"
GOLD = "A66A00"
BODY = "1F2937"
WHITE = "FFFFFF"
BORDER = "D5DDE8"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    """
    Đây là gì: hàm đặt kiểu chữ cho một "run" (một đoạn chữ nhỏ trong file Word có cùng một kiểu
      chữ).
    Mục tiêu: giúp chữ trong tài liệu dùng đúng phông, cỡ, màu trên mọi máy, kể cả chữ tiếng Việt.
    Chức năng: ghi tên phông (mặc định Calibri) cho cả ba nhóm ký tự ascii, hAnsi, eastAsia; nếu có
      truyền vào thì đặt thêm cỡ chữ (đơn vị pt), màu (mã màu như "0B2545"), in đậm, in nghiêng; thứ
      gì không truyền vào (để None) thì giữ nguyên.
    Đầu ra: không trả về gì; đoạn chữ run được sửa trực tiếp.
    Dùng ở đâu: add_cover, add_prompt_block, add_commands_and_presentation, add_page_field,
      populate_running_header, populate_running_footer (cùng file).
    Liên kết với: các hằng số màu đầu file (NAVY, MUTED, BODY...); thư viện python-docx (Pt,
      RGBColor, qn).
    """
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    """
    Đây là gì: hàm đặt lề bên trong một ô của bảng.
    Mục tiêu: để chữ trong ô không dính sát đường viền, bảng nhìn thoáng và dễ đọc.
    Chức năng: tìm phần cài đặt lề ô (thẻ XML "w:tcMar"), chưa có thì tạo mới; rồi ghi lề trên, lề
      đầu dòng, lề dưới, lề cuối dòng theo đơn vị dxa (mặc định 80, 120, 80, 120).
    Đầu ra: không trả về gì; ô bảng được sửa trực tiếp.
    Dùng ở đâu: configure_table (cùng file).
    Liên kết với: configure_table (gọi hàm này cho từng ô của bảng).
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    """
    Đây là gì: hàm tô màu nền cho một ô bảng.
    Mục tiêu: làm nổi bật hàng tiêu đề của bảng.
    Chức năng: tìm thẻ XML "w:shd" (tô nền) của ô, chưa có thì tạo mới, rồi ghi mã màu fill vào.
    Đầu ra: không trả về gì; ô bảng được sửa trực tiếp.
    Dùng ở đâu: configure_table (cùng file), để tô hàng tiêu đề.
    Liên kết với: configure_table; màu mặc định LIGHT_BLUE ở đầu file.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    """
    Đây là gì: hàm đánh dấu một hàng là "hàng tiêu đề lặp lại" của bảng.
    Mục tiêu: khi bảng dài tràn sang trang mới, Word tự in lại hàng tiêu đề ở đầu trang để người đọc
      không quên tên các cột.
    Chức năng: thêm thẻ XML "w:tblHeader" với giá trị true vào phần cài đặt của hàng.
    Đầu ra: không trả về gì; hàng được sửa trực tiếp.
    Dùng ở đâu: configure_table (cùng file).
    Liên kết với: configure_table (chỉ áp dụng cho hàng đầu tiên).
    """
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_table(table, widths_dxa, header_fill=LIGHT_BLUE):
    """
    Đây là gì: hàm định dạng toàn bộ một bảng trong file Word.
    Mục tiêu: cho mọi bảng trong tài liệu có cùng độ rộng, đường viền, lề ô và hàng tiêu đề màu
      giống nhau.
    Chức năng: kiểm tra tổng độ rộng các cột phải đúng 9360 dxa (bằng 6,5 inch: giấy rộng 8,5 inch
      trừ hai lề 1 inch; sai thì chương trình dừng và báo lỗi); căn bảng sang trái, tắt tự co giãn,
      dùng độ rộng cố định, thụt vào 120 dxa; kẻ viền mảnh màu BORDER cho mọi cạnh; đặt độ rộng từng
      cột; mỗi ô được căn giữa theo chiều dọc và đặt lề; hàng đầu tiên được tô màu header_fill và
      lặp lại khi sang trang.
    Đầu ra: không trả về gì; bảng được sửa trực tiếp.
    Dùng ở đâu: add_commands_and_presentation (cùng file), cho bảng câu lệnh Terminal.
    Liên kết với: set_cell_margins, shade_cell, set_repeat_table_header; configure_section_geometry
      (lề trang quyết định con số 9360).
    """
    assert sum(widths_dxa) == 9360
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(widths_dxa[col_index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[col_index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            if row_index == 0:
                shade_cell(cell, header_fill)
        if row_index == 0:
            set_repeat_table_header(row)


def add_numbering_definition(document, *, kind, marker):
    """
    Đây là gì: hàm tạo một kiểu danh sách mới trong file Word (gạch đầu dòng hoặc đánh số).
    Mục tiêu: có danh sách dấu chấm tròn "•" hoặc đánh số 1., 2., 3. thụt lề đều và đẹp.
    Chức năng: đọc các kiểu danh sách đã có trong tài liệu, chọn mã mới lớn hơn mã lớn nhất; tạo mô
      tả một cấp danh sách (bắt đầu từ 1, kiểu kind là "bullet" hoặc "decimal", ký hiệu marker, thụt
      lề 540, khoảng cách dòng); với gạch đầu dòng thì dùng phông Arial cho dấu chấm; chèn mô tả này
      vào trước thẻ "w:num" đầu tiên (Word đòi đúng thứ tự này) rồi thêm thẻ "w:num" trỏ tới nó.
    Đầu ra: số mã num_id của danh sách mới, dùng để gắn vào các đoạn văn.
    Dùng ở đâu: build_document (cùng file), tạo BULLET_ID, WORKFLOW_ID, ARCHITECTURE_ID.
    Liên kết với: assign_numbering (dùng mã num_id này); add_bullet, add_numbered.
    """
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), kind)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), marker)
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "271")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    if kind == "bullet":
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        r_pr.append(fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def assign_numbering(paragraph, num_id):
    """
    Đây là gì: hàm gắn một đoạn văn vào một danh sách đã tạo.
    Mục tiêu: biến đoạn văn thường thành một dòng có dấu chấm hoặc có số thứ tự.
    Chức năng: tìm phần "w:numPr" của đoạn văn, chưa có thì tạo; thêm cấp 0 (cấp ngoài cùng) và mã
      num_id.
    Đầu ra: không trả về gì; đoạn văn được sửa trực tiếp.
    Dùng ở đâu: add_bullet và add_numbered (cùng file).
    Liên kết với: add_numbering_definition (nơi tạo ra num_id).
    """
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_bullet(document, num_id, text, bold_prefix=None):
    """
    Đây là gì: hàm thêm một dòng gạch đầu dòng vào tài liệu.
    Mục tiêu: trình bày các ý ngắn thành danh sách dễ đọc.
    Chức năng: tạo đoạn văn mới, gắn vào danh sách num_id, đặt khoảng cách dòng; nếu có bold_prefix
      và câu bắt đầu đúng bằng phần đó (ví dụ "Kiểm tra webcam:") thì in đậm phần đầu, phần còn lại
      chữ thường; nếu không thì ghi cả câu chữ thường.
    Đầu ra: đoạn văn vừa tạo.
    Dùng ở đâu: add_intro (cùng file).
    Liên kết với: assign_numbering; biến toàn cục BULLET_ID (tạo trong build_document).
    """
    paragraph = document.add_paragraph()
    assign_numbering(paragraph, num_id)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_numbered(document, num_id, title, detail):
    """
    Đây là gì: hàm thêm một mục trong danh sách đánh số, dạng "Tiêu đề: chi tiết".
    Mục tiêu: trình bày các bước hoặc các tầng theo thứ tự 1, 2, 3.
    Chức năng: tạo đoạn văn mới, gắn vào danh sách num_id, đặt khoảng cách; ghi "title: " in đậm rồi
      ghi detail chữ thường.
    Đầu ra: đoạn văn vừa tạo.
    Dùng ở đâu: hiện không hàm nào trong file này gọi tới; được giữ lại vì bộ hàm phụ trợ được chép
      chung với build_ho_so_du_an.py và build_mo_ta_tinh_nang.py.
    Liên kết với: assign_numbering; biến toàn cục WORKFLOW_ID, ARCHITECTURE_ID (tạo trong
      build_document).
    """
    paragraph = document.add_paragraph()
    assign_numbering(paragraph, num_id)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    title_run = paragraph.add_run(f"{title}: ")
    title_run.bold = True
    paragraph.add_run(detail)
    return paragraph


def set_paragraph_shading(paragraph, fill, left_border=None):
    """
    Đây là gì: hàm tô nền cho cả một đoạn văn, có thể kèm vạch màu bên trái.
    Mục tiêu: tạo hiệu ứng "hộp" nổi bật cho ghi chú hoặc câu lệnh mẫu.
    Chức năng: tìm hoặc tạo thẻ XML "w:shd" rồi ghi màu nền fill; nếu có left_border thì thêm viền
      trái dày (cỡ 18), cách chữ 8, màu left_border.
    Đầu ra: không trả về gì; đoạn văn được sửa trực tiếp.
    Dùng ở đâu: add_callout và add_prompt_block (cùng file).
    Liên kết với: add_callout.
    """
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if left_border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        p_bdr.append(left)


def add_callout(document, label, text, fill=PALE_BLUE, accent=BLUE):
    """
    Đây là gì: hàm thêm một hộp ghi chú nổi bật vào tài liệu.
    Mục tiêu: làm người đọc chú ý tới nguyên tắc hoặc lưu ý quan trọng.
    Chức năng: tạo đoạn văn kiểu "Callout", tô nền màu fill và vạch trái màu accent; viết nhãn
      "label: " in đậm cùng màu accent, sau đó là nội dung text.
    Đầu ra: đoạn văn vừa tạo.
    Dùng ở đâu: add_intro và add_commands_and_presentation (cùng file).
    Liên kết với: set_paragraph_shading; configure_styles (nơi tạo kiểu chữ "Callout"); màu
      PALE_BLUE, BLUE.
    """
    paragraph = document.add_paragraph(style="Callout")
    set_paragraph_shading(paragraph, fill, accent)
    lead = paragraph.add_run(f"{label}: ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(accent)
    paragraph.add_run(text)
    return paragraph


def add_page_field(paragraph):
    """
    Đây là gì: hàm chèn số trang tự động vào một đoạn văn.
    Mục tiêu: mỗi trang tự hiện đúng số trang của nó mà không phải gõ tay.
    Chức năng: thêm một "trường" PAGE (ô đặc biệt mà Word tự điền số trang khi mở hoặc in) gồm ba
      phần: điểm bắt đầu, lệnh " PAGE ", điểm kết thúc; đặt chữ cỡ 9 màu MUTED.
    Đầu ra: không trả về gì; đoạn văn được thêm số trang.
    Dùng ở đâu: populate_running_footer (cùng file).
    Liên kết với: set_run_font.
    """
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(document):
    """
    Đây là gì: hàm cài đặt các kiểu chữ chung cho cả tài liệu.
    Mục tiêu: cho toàn bộ tài liệu thống nhất phông, cỡ chữ, màu và khoảng cách.
    Chức năng: kiểu Normal (chữ thường): Calibri 11, màu BODY, căn trái, giãn dòng 1,25; Title (tiêu
      đề): cỡ 30, đậm, màu NAVY, căn giữa, bỏ đường viền mặc định; Subtitle (phụ đề): cỡ 15, căn
      giữa; Heading 1, 2, 3 (đề mục): cỡ, màu, khoảng cách riêng và luôn dính với đoạn sau; tạo thêm
      kiểu "Lead" (đoạn mở đầu chữ to hơn) và "Callout" (hộp ghi chú).
    Đầu ra: không trả về gì; bảng kiểu chữ của tài liệu được sửa.
    Dùng ở đâu: build_document (cùng file).
    Liên kết với: add_callout (dùng kiểu Callout); các đoạn đặt style = "Lead"; các hằng số màu đầu
      file.
    """
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string("2B5163")
    subtitle.font.italic = False
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    lead = styles.add_style("Lead", 1)
    lead.base_style = normal
    lead.font.size = Pt(12.5)
    lead.font.color.rgb = RGBColor.from_string(NAVY)
    lead.paragraph_format.space_after = Pt(10)
    lead.paragraph_format.line_spacing = 1.25

    callout = styles.add_style("Callout", 1)
    callout.base_style = normal
    callout.font.size = Pt(10.5)
    callout.paragraph_format.left_indent = Inches(0.14)
    callout.paragraph_format.right_indent = Inches(0.08)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.2


def configure_section_geometry(section):
    """
    Đây là gì: hàm đặt khổ giấy và lề cho một phần (section) của tài liệu.
    Mục tiêu: cho mọi trang cùng kích thước giấy và lề.
    Chức năng: giấy Letter 8,5 x 11 inch, lề 1 inch mỗi bên, đầu trang và chân trang cách mép 0,492
      inch.
    Đầu ra: không trả về gì; section được sửa trực tiếp.
    Dùng ở đâu: configure_page và start_body_section (cùng file).
    Liên kết với: configure_table (độ rộng bảng 9360 dxa tính theo lề này).
    """
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_page(document):
    """
    Đây là gì: hàm chuẩn bị phần đầu tiên của tài liệu (phần trang bìa).
    Mục tiêu: trang bìa sạch, không có đầu trang hay số trang.
    Chức năng: bật chế độ đầu trang và chân trang khác nhau cho trang chẵn và trang lẻ; đặt khổ giấy
      cho phần đầu; xóa trống đầu trang, chân trang của phần này.
    Đầu ra: không trả về gì; tài liệu được sửa trực tiếp.
    Dùng ở đâu: build_document (cùng file).
    Liên kết với: configure_section_geometry; start_body_section (phần thân có đầu trang, chân trang
      riêng).
    """
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    configure_section_geometry(section)
    section.different_first_page_header_footer = False
    section.header.paragraphs[0].text = ""
    section.even_page_header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""
    section.even_page_footer.paragraphs[0].text = ""


def populate_running_header(header):
    """
    Đây là gì: hàm viết dòng đầu trang lặp lại ở mọi trang nội dung.
    Mục tiêu: người đọc luôn biết đang xem tài liệu nào.
    Chức năng: xóa chữ cũ, căn phải, ghi "AI EXAMGUARD  |  TÀI LIỆU MÔ TẢ HỆ THỐNG" chữ nhỏ cỡ 8,5,
      đậm, màu xám.
    Đầu ra: không trả về gì; đầu trang được sửa trực tiếp.
    Dùng ở đâu: start_body_section (cùng file), cho trang lẻ và trang chẵn.
    Liên kết với: set_run_font.
    """
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("AI EXAMGUARD  |  TÀI LIỆU MÔ TẢ HỆ THỐNG")
    set_run_font(run, size=8.5, color=MUTED, bold=True)


def populate_running_footer(footer):
    """
    Đây là gì: hàm viết dòng chân trang có số trang.
    Mục tiêu: mỗi trang nội dung có chữ "AI ExamGuard • Trang <số>" ở giữa chân trang.
    Chức năng: xóa chữ cũ, căn giữa, bỏ khoảng cách trên dưới, ghi chữ cỡ 9 màu xám rồi chèn số
      trang tự động.
    Đầu ra: không trả về gì; chân trang được sửa trực tiếp.
    Dùng ở đâu: start_body_section (cùng file), cho trang lẻ và trang chẵn.
    Liên kết với: add_page_field, set_run_font.
    """
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("AI ExamGuard  •  Trang ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(p)


def start_body_section(document):
    """
    Đây là gì: hàm bắt đầu phần thân của tài liệu ở một trang mới.
    Mục tiêu: tách trang bìa với các trang nội dung: bìa sạch, còn trang nội dung có đầu trang và số
      trang.
    Chức năng: thêm một "section" mới (một phần của tài liệu có thể có đầu trang, chân trang riêng)
      bắt đầu từ trang mới; đặt khổ giấy; ngắt liên kết đầu trang, chân trang với phần bìa; điền đầu
      trang và chân trang cho cả trang lẻ và trang chẵn.
    Đầu ra: không trả về gì; tài liệu có thêm phần thân.
    Dùng ở đâu: build_document (cùng file), ngay sau add_cover.
    Liên kết với: configure_section_geometry, populate_running_header, populate_running_footer.
    """
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section_geometry(section)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.even_page_header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.even_page_footer.is_linked_to_previous = False
    populate_running_header(section.header)
    populate_running_header(section.even_page_header)
    populate_running_footer(section.footer)
    populate_running_footer(section.even_page_footer)


def add_cover(document):
    """
    Đây là gì: hàm tạo trang bìa của tài liệu hướng dẫn.
    Mục tiêu: trang đầu cho biết ngay đây là tài liệu hướng dẫn và mẫu câu lệnh cho AI.
    Chức năng: thêm 2 dòng trống; dòng nhỏ màu vàng "TÀI LIỆU HƯỚNG DẪN KỸ THUẬT & MẪU PROMPT CHO
      AI"; tiêu đề "AI EXAMGUARD PYTHON"; phụ đề; khẩu hiệu in nghiêng; "Phiên bản tài liệu 1.0
      (Python Desktop)"; ngày 01 tháng 09 năm 2026.
    Đầu ra: không trả về gì; tài liệu có thêm trang bìa.
    Dùng ở đâu: build_document (cùng file), được gọi đầu tiên.
    Liên kết với: set_run_font; kiểu Title và Subtitle trong configure_styles.
    """
    for _ in range(2):
        document.add_paragraph().paragraph_format.space_after = Pt(14)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    set_run_font(kicker.add_run("TÀI LIỆU HƯỚNG DẪN KỸ THUẬT & MẪU PROMPT CHO AI"), size=11, color=GOLD, bold=True)
    document.add_paragraph("AI EXAMGUARD PYTHON", style="Title")
    document.add_paragraph("Hướng dẫn chi tiết từng bước xây dựng ứng dụng bằng OpenAI Codex", style="Subtitle")
    tagline = document.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline.paragraph_format.space_before = Pt(6)
    tagline.paragraph_format.space_after = Pt(40)
    set_run_font(tagline.add_run("Dành cho người không chuyên lập trình • Copy & Paste 5 bước • Đóng gói 1-Click .exe"), size=10.5, color=MUTED, italic=True)

    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(6)
    set_run_font(info.add_run("Phiên bản tài liệu 1.0 (Python Desktop)"), size=10.5, color=NAVY, bold=True)
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_after = Pt(4)
    set_run_font(date.add_run("Ngày 01 tháng 09 năm 2026"), size=10, color=MUTED)


def add_intro(document):
    """
    Đây là gì: hàm viết mục 1 "Tổng quan & Mục đích tài liệu".
    Mục tiêu: giải thích tài liệu dùng để làm gì và kiến trúc ứng dụng sẽ được tạo ra.
    Chức năng: thêm đoạn mở đầu kiểu Lead; hộp xanh "Nguyên tắc thực hiện" (gửi lần lượt 5 bước cho
      AI); mục 1.1 kiến trúc với 4 gạch đầu dòng: chương trình chạy main.py, máy chủ web nội bộ, AI
      MediaPipe chạy không cần mạng, đóng gói thành một file .exe.
    Đầu ra: không trả về gì; tài liệu có thêm mục 1.
    Dùng ở đâu: build_document (cùng file), ngay sau start_body_section.
    Liên kết với: add_callout, add_bullet, biến toàn cục BULLET_ID.
    """
    document.add_heading("1. Tổng quan & Mục đích tài liệu", level=1)
    document.add_paragraph(
        "Tài liệu này được biên soạn nhằm cung cấp toàn bộ quy trình, kiến trúc kỹ thuật và đặc biệt là 5 mẫu câu lệnh (Prompt) chuẩn "
        "để bạn chỉ cần sao chép và dán trực tiếp vào OpenAI Codex hoặc trợ lý AI nhằm tự động xây dựng hệ thống giám sát thi AI ExamGuard bằng Python."
    ).style = "Lead"
    add_callout(
        document,
        "Nguyên tắc thực hiện",
        "Bạn không cần biết lập trình. Hãy mở một khung chat mới với Codex/AI, gửi lần lượt từ Bước 1 đến Bước 5. AI sẽ tự động sinh mã nguồn, tải tài nguyên AI ngoại tuyến và đóng gói thành file .exe hoàn chỉnh.",
        fill=PALE_GREEN,
        accent=GREEN,
    )
    document.add_heading("1.1. Kiến trúc hệ thống Python Desktop", level=2)
    add_bullet(document, BULLET_ID, "Python App Runner (main.py): Khởi tạo cửa sổ Desktop chuẩn Windows bằng Microsoft Edge WebView2 qua thư viện pywebview, tự động cấp cổng localhost và quản lý tiến trình.", "Python App Runner:")
    add_bullet(document, BULLET_ID, "Local HTTP Server: Web server nội bộ đa luồng hỗ trợ Clean URL (/thi-sinh, /giam-thi, /quan-tri) và phân giải đúng MIME type cho .wasm, .js, .css.", "Local HTTP Server:")
    add_bullet(document, BULLET_ID, "Edge AI Offline 100%: Mô hình Google MediaPipe Face Landmarker lưu trữ trực tiếp trong thư mục ui/mediapipe/ giúp máy không có mạng Internet vẫn chạy AI 60 FPS.", "Edge AI Offline:")
    add_bullet(document, BULLET_ID, "Đóng gói 1-Click .exe: Sử dụng PyInstaller để tạo ra đúng 1 file duy nhất AI_ExamGuard_Portable.exe cắm USB sang máy khác chạy ngay.", "Đóng gói 1-Click .exe:")


def add_prompt_block(document, step_num, title, prompt_text):
    """
    Đây là gì: hàm thêm một bước gồm tiêu đề và khung chứa câu lệnh mẫu.
    Mục tiêu: người đọc thấy rõ phần nào cần sao chép nguyên văn để dán vào Codex.
    Chức năng: thêm đề mục "Bước <số>: <tiêu đề>"; dòng hướng dẫn "Câu lệnh (Prompt) sao chép và dán
      vào Codex / AI"; rồi một khung nền xám có vạch xanh bên trái chứa câu lệnh, viết bằng phông
      Consolas (phông chữ các ký tự rộng đều nhau, kiểu chữ máy tính).
    Đầu ra: không trả về gì; tài liệu có thêm một bước.
    Dùng ở đâu: add_steps_for_codex (cùng file), gọi 5 lần cho 5 bước.
    Liên kết với: set_paragraph_shading, set_run_font.
    """
    document.add_heading(f"Bước {step_num}: {title}", level=2)
    p_desc = document.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(3)
    r_desc = p_desc.add_run("👉 Câu lệnh (Prompt) sao chép và dán vào Codex / AI:")
    set_run_font(r_desc, size=10.5, color=DARK_BLUE, bold=True)

    p_box = document.add_paragraph()
    set_paragraph_shading(p_box, "F4F6F8", BLUE)
    p_box.paragraph_format.left_indent = Inches(0.18)
    p_box.paragraph_format.right_indent = Inches(0.1)
    p_box.paragraph_format.space_before = Pt(3)
    p_box.paragraph_format.space_after = Pt(12)
    p_box.paragraph_format.line_spacing = 1.15
    r_text = p_box.add_run(prompt_text.strip())
    set_run_font(r_text, name="Consolas", size=9.5, color="0A2540")


def add_steps_for_codex(document):
    """
    Đây là gì: hàm viết mục 2 "Lộ trình 5 bước giao việc cho Codex / AI".
    Mục tiêu: cung cấp 5 câu lệnh mẫu (prompt: lời yêu cầu viết bằng tiếng thường để giao việc cho
      trợ lý AI lập trình Codex) theo đúng thứ tự.
    Chức năng: soạn sẵn 5 câu lệnh: 1 tạo thư mục và danh sách thư viện; 2 viết main.py mở cửa sổ và
      máy chủ web nội bộ; 3 tải mô hình AI MediaPipe về máy; 4 làm 4 màn hình giao diện trong ui/; 5
      tạo file .bat để chạy và đóng gói .exe; mỗi câu lệnh được thêm bằng add_prompt_block.
    Đầu ra: không trả về gì; tài liệu có thêm mục 2.
    Dùng ở đâu: build_document (cùng file), sau add_intro.
    Liên kết với: add_prompt_block. Kiến trúc trong các câu lệnh (main.py, ui/) là bản cũ, khác mã
      nguồn hiện tại ở web/ và desktop-hub/.
    """
    document.add_heading("2. Lộ trình 5 bước giao việc cho Codex / AI", level=1)

    step1 = (
        "Tôi muốn bạn tạo một dự án Python Desktop có tên là \"AI_ExamGuard\" (Hệ thống giám sát thi thông minh bằng AI) có thể chạy trên Windows và đóng gói thành file .exe độc lập.\n\n"
        "Trước tiên, hãy:\n"
        "1. Tạo cấu trúc thư mục dự án gồm:\n"
        "   - Thư mục gốc: chứa main.py, requirements.txt, các file .bat khởi chạy\n"
        "   - Thư mục ui/: chứa toàn bộ giao diện HTML, CSS, Javascript và mô hình AI\n"
        "   - Thư mục ui/mediapipe/: chứa tài nguyên mô hình AI ngoại tuyến (Offline)\n"
        "   - Thư mục dist/: chứa file .exe sau khi đóng gói\n"
        "2. Tạo file requirements.txt chứa các thư viện:\n"
        "   pywebview>=6.2.0\n"
        "   pythonnet>=3.0.0\n"
        "   pyinstaller>=6.20.0\n"
        "   python-docx>=1.2.0\n"
        "3. Hướng dẫn tôi câu lệnh Terminal để cài đặt các thư viện này."
    )
    add_prompt_block(document, 1, "Khởi tạo cấu trúc thư mục & Danh sách thư viện", step1)

    step2 = (
        "Bây giờ, hãy viết file mã nguồn main.py hoàn chỉnh để chạy ứng dụng Desktop:\n\n"
        "Yêu cầu kỹ thuật trong file main.py:\n"
        "1. Sử dụng thư viện pywebview để mở cửa sổ Desktop chuẩn Windows (dùng engine Edge WebView2 với độ phân giải 1380x880).\n"
        "2. Tích hợp một Web Server cục bộ đa luồng (http.server.HTTPServer) chạy trên cổng localhost ngẫu nhiên đang rảnh (127.0.0.1:<free_port>).\n"
        "3. Web Server phải phục vụ thư mục ui/, hỗ trợ Clean URL (tự động mở /thi-sinh -> ui/thi-sinh.html, /giam-thi -> ui/giam-thi.html), định dạng đúng MIME type cho .wasm, .js, .css, .json, và không lưu cache.\n"
        "4. Viết hàm get_resource_path() thông minh để khi đóng gói thành file .exe qua PyInstaller (dùng sys._MEIPASS hoặc _internal) vẫn load được toàn bộ giao diện và mô hình AI.\n"
        "5. Tự động dọn dẹp và tắt sạch tiến trình ngầm khi người dùng đóng cửa sổ ứng dụng."
    )
    add_prompt_block(document, 2, "Viết mã nguồn Python Desktop & Web Server nội bộ (main.py)", step2)

    step3 = (
        "Hãy viết cho tôi một script Python nhỏ download_ai_models.py và chạy nó để tự động tải toàn bộ tài nguyên AI Google MediaPipe về lưu cục bộ vào thư mục ui/mediapipe/ (giúp ứng dụng chạy được ngay cả khi máy không có mạng Internet):\n\n"
        "1. Tải 6 file WASM từ npm @mediapipe/tasks-vision/wasm lưu vào ui/mediapipe/wasm/.\n"
        "2. Tải file mô hình nhận diện khuôn mặt 478 điểm mốc:\n"
        "   https://storage.googleapis.com/mediapipe-models/face_landmarker/face_landmarker/float16/1/face_landmarker.task\n"
        "   lưu vào ui/mediapipe/face_landmarker.task."
    )
    add_prompt_block(document, 3, "Tự động tải mô hình AI Google MediaPipe để chạy Offline", step3)

    step4 = (
        "Bây giờ hãy tạo đầy đủ các file giao diện HTML/CSS/Javascript hiện đại (giao diện tối Dark Mode công nghệ cao, dùng Tailwind CSS, icon Lucide và Canvas AI) trong thư mục ui/:\n\n"
        "1. File ui/index.html (Trang chủ):\n"
        "   - Giới thiệu hệ thống AI ExamGuard.\n"
        "   - Nút bấm điều hướng vào 3 phân hệ: Thí sinh, Giám thị, Quản trị.\n\n"
        "2. File ui/thi-sinh.html (Phân hệ Thí sinh):\n"
        "   - Form khai báo họ tên, số máy (PC-01, PC-02...), phòng thi và cam kết quy chế.\n"
        "   - Luồng video Webcam hiển thị trực tiếp.\n"
        "   - Chạy mô hình MediaPipe Face Landmarker từ file cục bộ ui/mediapipe/ để vẽ khung xanh/đỏ quanh khuôn mặt và các điểm mốc mắt/mũi.\n"
        "   - Thuật toán Behavior Engine tính toán: góc quay đầu (Yaw), góc cúi/ngẩng (Pitch), hướng mắt (Gaze), phát hiện vắng mặt hoặc có người thứ hai.\n"
        "   - Bắt sự kiện chuyển tab và phím tắt cấm (Ctrl+C, Ctrl+V, PrintScreen).\n"
        "   - Đồng bộ sự kiện qua BroadcastChannel sang màn hình Giám thị.\n\n"
        "3. File ui/giam-thi.html (Dashboard Giám thị):\n"
        "   - Sơ đồ phòng thi với các ô máy thi đổi màu theo thời gian thực (Xanh: Ổn định, Vàng: Cảnh báo, Đỏ: Nguy cơ).\n"
        "   - Luồng sự kiện cảnh báo trực tiếp (Live Alert Stream) có bộ lọc và tìm kiếm.\n"
        "   - Bảng xem chi tiết vi phạm và nút \"Đánh dấu đã xác minh\".\n"
        "   - Nút \"Xuất báo cáo\" tải file CSV danh sách vi phạm của ca thi.\n\n"
        "4. File ui/quan-tri.html (Cấu hình hệ thống):\n"
        "   - Cho phép chỉnh sửa ngưỡng thời gian cảnh báo (số giây vắng mặt, quay đầu, lệch mắt, cooldown).\n"
        "   - Tùy chọn quyền riêng tư và xuất cấu hình JSON."
    )
    add_prompt_block(document, 4, "Thiết kế toàn bộ 4 màn hình giao diện (ui/)", step4)

    step5 = (
        "Hãy tạo cho tôi 2 file batch script trên Windows để tôi có thể sử dụng chỉ bằng 1 cú nhấp đúp chuột:\n\n"
        "1. File chay_ung_dung.bat:\n"
        "   - Lệnh tự động chạy python main.py để tôi xem trước ứng dụng.\n\n"
        "2. File dong_goi_exe.bat:\n"
        "   - Lệnh tự động chạy PyInstaller:\n"
        "     pyinstaller --onefile --noconsole --name \"AI_ExamGuard_Portable\" --add-data \"ui;ui\" --clean -y main.py\n"
        "   - Đóng gói toàn bộ code Python, giao diện và mô hình AI thành đúng 1 file duy nhất dist/AI_ExamGuard_Portable.exe.\n\n"
        "Sau đó hãy chạy thử file dong_goi_exe.bat và báo cho tôi vị trí file .exe đã tạo thành công."
    )
    add_prompt_block(document, 5, "Tạo file khởi chạy & Đóng gói file .exe 1-Click", step5)


def add_commands_and_presentation(document):
    """
    Đây là gì: hàm viết mục 3 "Bảng tổng hợp các câu lệnh Terminal chính" và mục 4 "Kinh nghiệm
      thuyết trình khi đi thi".
    Mục tiêu: gom các lệnh hay dùng vào một bảng để tra nhanh, kèm lời khuyên khi trình bày trước
      ban giám khảo.
    Chức năng: tạo bảng 3 cột (Mục đích, Câu lệnh thực thi, Ghi chú) với 4 hàng: cài thư viện, chạy
      ứng dụng, đóng gói 1 file .exe, đóng gói dạng thư mục; cột câu lệnh dùng phông Consolas; độ
      rộng cột 1800/4760/2800 dxa; sau đó thêm hộp "Lời khuyên thuyết trình" nền vàng nhạt.
    Đầu ra: không trả về gì; tài liệu có thêm mục 3 và 4.
    Dùng ở đâu: build_document (cùng file), được gọi cuối cùng trước khi lưu.
    Liên kết với: set_run_font, configure_table, add_callout.
    """
    document.add_heading("3. Bảng tổng hợp các câu lệnh Terminal chính", level=1)
    table = document.add_table(rows=1, cols=3)
    headers = ["Mục đích", "Câu lệnh thực thi", "Ghi chú"]
    for index, text in enumerate(headers):
        p = table.rows[0].cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=10, color=NAVY, bold=True)
    cmd_rows = [
        ("Cài đặt thư viện", "pip install -r requirements.txt", "Cài đặt pywebview, pyinstaller, pythonnet..."),
        ("Chạy ứng dụng Python", "python main.py", "Mở cửa sổ Desktop kiểm tra trực tiếp"),
        ("Đóng gói 1 file .exe", "pyinstaller --onefile --noconsole --name AI_ExamGuard_Portable --add-data \"ui;ui\" --clean -y main.py", "Tạo file dist/AI_ExamGuard_Portable.exe gửi máy khác"),
        ("Đóng gói dạng thư mục", "pyinstaller --noconsole --name AI_ExamGuard --add-data \"ui;ui\" --clean -y main.py", "Tạo thư mục dist/AI_ExamGuard/ mở tức thì"),
    ]
    for purpose, command, note in cmd_rows:
        cells = table.add_row().cells
        for index, text in enumerate((purpose, command, note)):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index == 0:
                set_run_font(p.add_run(text), size=9.7, color=DARK_BLUE, bold=True)
            elif index == 1:
                set_run_font(p.add_run(text), name="Consolas", size=9, color=NAVY, bold=True)
            else:
                set_run_font(p.add_run(text), size=9.5, color=BODY)
    configure_table(table, [1800, 4760, 2800])

    document.add_heading("4. Kinh nghiệm thuyết trình khi đi thi", level=1)
    add_callout(
        document,
        "Lời khuyên thuyết trình",
        "Khi ban giám khảo hỏi về kỹ thuật, hãy nhấn mạnh 3 điểm chính: 1) Toàn bộ AI chạy trực tiếp trên máy thí sinh (Edge AI) bảo vệ quyền riêng tư; 2) Hệ thống hoạt động 100% Offline không cần mạng; 3) Ứng dụng đã được đóng gói thành file .exe tiêu chuẩn Windows có thể cắm USB chạy ngay.",
        fill="FFF8E8",
        accent=GOLD,
    )


def build_document():
    """
    Đây là gì: hàm chính, ghép mọi phần lại thành file Word hoàn chỉnh rồi lưu ra đĩa.
    Mục tiêu: chỉ cần gọi một hàm là tạo xong cả tài liệu.
    Chức năng: tạo tài liệu trống; ghi thông tin tài liệu (tiêu đề, chủ đề, tác giả); cài kiểu chữ
      và trang; tạo 3 kiểu danh sách và cất mã vào biến toàn cục BULLET_ID, WORKFLOW_ID,
      ARCHITECTURE_ID (trong file này chỉ BULLET_ID thực sự được dùng); thêm trang bìa, phần thân và
      lần lượt add_intro, add_steps_for_codex, add_commands_and_presentation; tạo thư mục nếu chưa
      có rồi lưu file và in đường dẫn ra màn hình.
    Đầu ra: cùng một file Word được lưu ra hai nơi OUTPUT_1 và OUTPUT_2; không trả về giá trị.
    Dùng ở đâu: khối if __name__ == "__main__" ở cuối file, tức là khi chạy tay lệnh python3
      docs/tools/build_huong_dan_codex.py.
    Liên kết với: configure_styles, configure_page, add_numbering_definition, add_cover,
      start_body_section và các hàm viết từng mục ở trên.
    """
    global BULLET_ID, WORKFLOW_ID, ARCHITECTURE_ID
    document = Document()
    document.core_properties.title = "AI ExamGuard Python - Hướng dẫn xây dựng bằng Codex"
    document.core_properties.subject = "Lộ trình 5 bước giao việc cho Codex xây dựng ứng dụng Python Desktop"
    document.core_properties.author = "AI ExamGuard Team"
    configure_styles(document)
    configure_page(document)
    BULLET_ID = add_numbering_definition(document, kind="bullet", marker="•")
    WORKFLOW_ID = add_numbering_definition(document, kind="decimal", marker="%1.")
    ARCHITECTURE_ID = add_numbering_definition(document, kind="decimal", marker="%1.")

    add_cover(document)
    start_body_section(document)
    add_intro(document)
    add_steps_for_codex(document)
    add_commands_and_presentation(document)

    OUTPUT_1.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_2.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_1)
    document.save(OUTPUT_2)
    print("Exported successfully to:", OUTPUT_1, "and", OUTPUT_2)


if __name__ == "__main__":
    build_document()

