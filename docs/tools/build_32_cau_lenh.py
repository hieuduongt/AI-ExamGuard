from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / 'tai-lieu' / 'AI_ExamGuard_32_Cau_Lenh_Tao_Du_An_Va_Build_EXE.docx'

# Each numbered step contains exactly one copyable instruction paragraph.
PAGES = [
    ('Khởi tạo và xác định phạm vi', [
        ('Thiết lập cách làm việc', 'Hãy tạo mới dự án AI ExamGuard trong thư mục hiện tại để hỗ trợ giám sát thi trên Windows qua mạng LAN, chỉ tạo cảnh báo nghi vấn để giáo viên xác minh, dùng Next.js 16, React 19, TypeScript, Tailwind CSS 4, MediaPipe Tasks Vision và Python với pywebview, làm lần lượt theo từng yêu cầu tôi gửi, thực hiện mã nguồn thật rồi kiểm tra, báo ngắn gọn tệp đã tạo và kết quả kiểm tra sau mỗi bước, không xóa hoặc ghi đè dữ liệu có sẵn và không tự triển khai lên Internet.'),
        ('Kiểm tra môi trường phát triển', 'Hãy kiểm tra thư mục dự án, Node.js 24, npm, Python 3.12, Git và khả năng dùng Microsoft Edge WebView2 trên Windows, tạo môi trường Python riêng cho dự án, nêu rõ thành phần còn thiếu và hướng dẫn tôi cài nếu cần quyền quản trị, đồng thời lưu các phiên bản phụ thuộc tương thích vào tệp khóa để có thể build lại.'),
        ('Tạo bộ khung dự án', 'Hãy khởi tạo giao diện Next.js App Router dùng TypeScript, Tailwind CSS và các thành phần shadcn phù hợp, tạo bốn trang /, /thi-sinh, /giam-thi và /quan-tri cùng các thư mục components, lib, tests, public/mediapipe, cấu hình xuất web tĩnh vào out, bổ sung lệnh lint, kiểm tra kiểu, test và build, rồi xác nhận bộ khung chạy được mà không yêu cầu Firebase.'),
    ]),
    ('Giao diện và dữ liệu dùng chung', [
        ('Xây dựng giao diện tiếng Việt', 'Hãy thiết kế giao diện AI ExamGuard tiếng Việt với nền sáng, thẻ thông tin dễ đọc, màu xanh cho trạng thái bình thường, vàng cho nghi vấn và đỏ cho ưu tiên kiểm tra, tạo thanh chuyển khu vực và trang chủ chọn Thí sinh hoặc Giám thị, bảo đảm dùng được trên màn hình 1366 × 768 và không hiển thị số liệu độ chính xác hoặc trạng thái AI giả.'),
        ('Định nghĩa dữ liệu phiên thi', 'Hãy tạo lib/exam-types.ts định nghĩa VisionMetrics, BehaviorThresholds, ExamEvent và StationStatus với mã phòng, số máy, họ tên, số báo danh, mã sự kiện duy nhất, loại cảnh báo, mức độ, thời điểm, trạng thái xác minh, nguồn ai hoặc browser hoặc demo, trạng thái phân tích và lần phân tích gần nhất, đồng thời thống nhất tên hiển thị tiếng Việt cho mọi loại cảnh báo.'),
        ('Tạo máy chủ LAN', 'Hãy tạo app_desktop.py dùng ThreadingHTTPServer phục vụ thư mục out và các API GET /api/health, GET /api/runtime, GET /api/snapshot, POST /api/stations, POST /api/events và PATCH /api/events/{id}, trả địa chỉ LAN cùng phiên bản ứng dụng từ runtime, dùng cổng 8765 mặc định và thông báo cổng thực tế nếu cổng này bận, lưu dữ liệu trong bộ nhớ cho bản LAN ban đầu và ghi rõ dữ liệu sẽ mất khi đóng máy chủ.'),
        ('Kiểm tra an toàn API và đường dẫn', 'Hãy kiểm tra dữ liệu đầu vào của các API để từ chối bản ghi thiếu danh tính hoặc sai kiểu, giới hạn kích thước yêu cầu và số sự kiện, xử lý đồng thời bằng khóa phù hợp, cho phép CORS cho kết nối LAN cần thiết, ngăn truy cập tệp ngoài out, ưu tiên thi-sinh.html và giam-thi.html thay vì trả danh sách thư mục cùng tên của Next.js, đồng thời cảnh báo bản LAN chưa có xác thực không được mở ra Internet.'),
        ('Tạo biểu mẫu thí sinh', 'Hãy xây dựng trang Thí sinh có họ tên, số báo danh, mã phòng, số máy hoặc vị trí, địa chỉ máy giám thị, nút Kiểm tra kết nối và xác nhận đã đọc thông báo giám sát, chuẩn hóa dữ liệu nhập và chặn bắt đầu khi thiếu thông tin hoặc chưa kết nối thành công, đồng thời nói rõ danh tính này do người dùng khai báo và cần giáo viên đối chiếu chứ không phải nhận diện danh tính bằng khuôn mặt.'),
    ]),
    ('Webcam và bộ phát hiện hành vi', [
        ('Xây dựng vòng đời webcam', 'Hãy triển khai xin quyền webcam bằng getUserMedia, hiển thị video phản chiếu và nút bắt đầu hoặc dừng, xử lý rõ lỗi từ chối quyền, không tìm thấy camera và camera bị ứng dụng khác chiếm, ngăn khởi động lặp và giải phóng track khi dừng, chỉ bật camera sau khi người dùng xác nhận và không thu âm hoặc quay bí mật.'),
        ('Đóng gói mô hình AI tại máy', 'Hãy tích hợp MediaPipe FaceLandmarker phát hiện tối đa hai khuôn mặt, tải mô hình face_landmarker.task và tệp WASM từ nguồn chính thức phù hợp phiên bản thư viện vào public/mediapipe, ưu tiên tài nguyên cục bộ và thử CPU nếu GPU khởi tạo thất bại, hiển thị riêng trạng thái đang tải, sẵn sàng và lỗi AI, đồng thời không báo AI đang hoạt động chỉ vì camera có hình.'),
        ('Tính các tín hiệu khuôn mặt', 'Hãy tách hàm tính số khuôn mặt, chỉ số quay đầu trái phải, cúi ngẩng, hướng mắt và FPS từ kết quả landmarks để có thể kiểm thử độc lập, vẽ khung mặt trên canvas, bảo vệ phép tính trước dữ liệu thiếu hoặc không hợp lệ, ghi rõ góc chỉ là ước tính cần hiệu chỉnh và không dùng FaceLandmarker để tuyên bố phát hiện điện thoại.'),
        ('Tạo bộ đếm hành vi có lọc nhiễu', 'Hãy tạo lib/behavior-engine.ts với ngưỡng mặc định không thấy mặt 5 giây, nhiều khuôn mặt 2 giây, quay đầu hoặc cúi ngẩng 3 giây, hướng mắt lệch 5 giây và khoảng nghỉ cảnh báo 12 giây, sử dụng ngưỡng tín hiệu ban đầu yaw 14, pitch 20, gazeX 0.55 và gazeY 0.7, giữ bộ đếm qua nhiễu ngắn dưới 400 mili giây nhưng không cộng thời gian nhiễu và xóa bộ đếm khi khung hình gián đoạn trên 1500 mili giây.'),
        ('Phân biệt nghi vấn với hành vi bình thường', 'Hãy bổ sung kiểm tra để một lần liếc nhanh không bị cộng dồn thành quay đầu kéo dài, chỉ đánh giá đầu và mắt khi có đúng một khuôn mặt, dùng cảnh báo mất mặt khi quay hẳn khỏi camera, mô tả cúi đầu lâu là cần kiểm tra bối cảnh vì học sinh có thể viết nháp và tuyệt đối không tự kết luận gian lận hoặc sử dụng điện thoại từ tư thế đầu.'),
    ]),
    ('Duy trì phân tích và gửi cảnh báo', [
        ('Giám sát tình trạng phân tích', 'Hãy chạy nhận diện trên các khung hình video thực sự mới bằng vòng lặp có giới hạn tần suất, không phụ thuộc hoàn toàn vào requestAnimationFrame, gửi nhịp trạng thái mỗi 2 giây với thời điểm phân tích cuối, báo AI tạm dừng khi quá 3 giây không có kết quả mới và báo máy mất kết nối khi quá 7 giây không có nhịp trạng thái, đồng thời không cộng thời gian camera đóng băng thành hành vi bất thường.'),
        ('Hiển thị tiến trình cảnh báo', 'Hãy hiển thị trên trang Thí sinh các tín hiệu thực tế và tiến trình từng cảnh báo theo dạng số giây quan sát trên số giây cần thiết, hiện rõ ngưỡng đang áp dụng, xóa tiến trình cũ khi phiên mới bắt đầu hoặc AI ngừng phân tích, bổ sung thông báo phân biệt chưa vượt ngưỡng, đang chờ đủ thời gian và đã tạo sự kiện để người dùng biết vì sao quay mặt chưa có cảnh báo.'),
        ('Phát hiện thao tác trình duyệt', 'Hãy tạo sự kiện khi trang thi bị ẩn và khi người dùng thực hiện các tổ hợp Ctrl hoặc Cmd với C, V, X hoặc P trong cửa sổ ứng dụng, gắn đúng nguồn browser và khoảng nghỉ chống gửi trùng, chỉ mô tả hành vi quan sát được và ghi rõ cơ chế này không theo dõi toàn bộ ứng dụng khác trên Windows.'),
        ('Kết nối thí sinh với máy giám thị', 'Hãy tạo lib/exam-store.ts chuẩn hóa và lưu địa chỉ giám thị do người dùng nhập, kiểm tra /api/health với thời gian chờ hữu hạn, gửi trạng thái và sự kiện bằng HTTP đến đúng hub đã chọn, đưa đủ phòng, máy, họ tên và số báo danh vào mọi bản ghi, không hardcode IP và không dùng riêng localStorage hoặc BroadcastChannel để giả làm đồng bộ giữa hai máy khác nhau.'),
        ('Bảo đảm cảnh báo được gửi lại', 'Hãy lưu hàng đợi cảnh báo chưa gửi trong localStorage kèm đúng địa chỉ hub đích và ID sự kiện, thử gửi lại khi có nhịp kết nối hoặc kiểm tra kết nối thành công, chỉ ghi Giám thị đã nhận sau phản hồi thành công của máy chủ, giữ cảnh báo khi lỗi mạng và làm POST sự kiện có tính lặp an toàn để gửi lại không nhân đôi hoặc xóa trạng thái đã xác minh.'),
    ]),
    ('Giám thị và cấu hình', [
        ('Hiển thị dữ liệu trên máy giáo viên', 'Hãy xây dựng trang Giám thị đọc /api/snapshot khoảng mỗi giây và chỉ có một yêu cầu đang chờ, hiển thị sơ đồ máy theo phòng với tên học sinh, số báo danh, số máy, trạng thái mạng và trạng thái AI, dùng dữ liệu máy chủ làm nguồn chính sau khi kết nối và không để dữ liệu minh họa cục bộ ghi đè dữ liệu thật.'),
        ('Thông báo rõ máy và học sinh', 'Hãy thêm dải Cảnh báo mới cần kiểm tra khi nhận sự kiện mới chưa xác minh, hiển thị nổi bật phòng, máy, họ tên, số báo danh, loại hành vi, thời điểm và mô tả thời gian quan sát, giữ sự kiện trong danh sách sau khi học sinh nhìn lại màn hình, đồng thời cung cấp bộ lọc theo mức độ, chưa xác minh và tìm theo danh tính.'),
        ('Xác minh và xuất báo cáo', 'Hãy bổ sung khung chi tiết cảnh báo, thao tác đánh dấu đã xác minh được lưu vào máy chủ và xuất báo cáo CSV tiếng Việt có BOM cùng xử lý ký tự đặc biệt an toàn, phân biệt sự kiện AI, trình duyệt và dữ liệu thử, đồng thời nhắc giáo viên xuất báo cáo trước khi đóng hub đang lưu dữ liệu trong bộ nhớ.'),
        ('Tạo trang quản trị cấu hình', 'Hãy tạo trang Quản trị để điều chỉnh ngưỡng thời gian, khoảng nghỉ cảnh báo và chế độ máy yếu, kiểm tra giá trị số hợp lệ, khôi phục mặc định và nhập xuất cấu hình JSON, ghi rõ cấu hình cục bộ chỉ có hiệu lực trên máy đã áp dụng và phiên tiếp theo, không tuyên bố thay đổi ở máy giáo viên tự động cập nhật mọi máy thí sinh nếu chưa triển khai đồng bộ cấu hình.'),
        ('Sửa lỗi đổi vai trò làm dừng camera', 'Hãy tạo nút Mở Giám thị ở trang Thí sinh để mở cửa sổ riêng và giữ nguyên phiên webcam, dùng js_api.open_proctor của pywebview trong bản desktop và cửa sổ mới khi chạy trình duyệt, truyền đúng hub đã chọn sang cửa sổ giám thị, chặn điều hướng gây mất phiên khi đang giám sát và hiển thị lỗi mở cửa sổ thay vì âm thầm dừng camera.'),
    ]),
    ('Kiểm thử trước khi đóng gói', [
        ('Kiểm thử bộ phát hiện tự động', 'Hãy viết và chạy unit test cho quay trái và phải đủ thời gian, quay đầu có khung hình nhiễu ngắn, liếc nhanh rồi nhìn thẳng, không có mặt, nhiều mặt, dừng rồi bắt đầu lại, khung hình đóng băng và khoảng nghỉ cảnh báo, kiểm tra cả trường hợp phải báo và không được báo, rồi sửa nguyên nhân nếu test thất bại chứ không chỉ hạ ngưỡng cho qua test.'),
        ('Kiểm thử đầy đủ luồng cảnh báo', 'Hãy tạo hai hub kiểm thử riêng và bài kiểm thử trình duyệt với hai ngữ cảnh Thí sinh và Giám thị, kiểm tra MediaPipe thật trên camera giả lập không có người tạo cảnh báo mất mặt, dữ liệu quay đầu mô phỏng đi qua bộ phát hiện rồi đến giao diện giám thị, đúng danh tính, hàng đợi khi lỗi mạng, gửi lại không trùng, trạng thái AI đóng băng và cửa sổ riêng không dừng webcam, đồng thời ghi rõ dữ liệu mô phỏng không thay thế kiểm tra webcam thật.'),
        ('Kiểm tra bằng webcam thực tế', 'Hãy mở ứng dụng để tôi cho phép webcam và phối hợp kiểm tra nhìn thẳng, quay mặt trái phải đủ thời gian, quay hẳn khỏi camera, cúi xuống viết nháp và ánh sáng yếu, ghi lại chỉ số, bộ đếm, trạng thái AI, xác nhận gửi và màn hình giáo viên cho từng ca, chỉ kết luận phần đã quan sát thực tế và giữ các ca chưa thử ở trạng thái chưa xác minh.'),
        ('Hoàn thiện ứng dụng desktop', 'Hãy hoàn thiện app_desktop.py để mở webview EdgeChromium tại địa chỉ loopback của máy hiện tại, phục vụ hub LAN trên giao diện mạng thích hợp, tìm tài nguyên đúng cả khi chạy mã nguồn và trong PyInstaller, mở giám thị bằng cửa sổ riêng có thể dùng lại, dọn tài nguyên khi đóng và ghi phiên bản trên tiêu đề cửa sổ cùng /api/runtime.'),
        ('Chuẩn bị bản phát hành', 'Hãy chạy lint, kiểm tra TypeScript, toàn bộ test và npm run build, xác nhận out có đầy đủ trang HTML, JavaScript, CSS, mô hình và WASM cục bộ, chạy thử bằng máy chủ Python với đường dẫn trực tiếp và tải lại từng trang, sửa hết lỗi ảnh hưởng vận hành rồi ghi kết quả đạt, chưa đạt hoặc chưa kiểm thử vào báo cáo kiểm tra trước phát hành.'),
    ]),
    ('Build EXE và bàn giao', [
        ('Tạo quy trình build EXE', 'Hãy tạo requirements.txt khóa phiên bản Python cần thiết, tệp AI_ExamGuard_Portable.spec dùng PyInstaller với chế độ một tệp và không mở console, đóng gói toàn bộ out, các thư viện pywebview cùng tài nguyên AI cần dùng, rồi tạo build_exe.ps1 dừng ngay khi lệnh lỗi và tự chạy build web trước khi đóng gói để tránh EXE chứa giao diện cũ.'),
        ('Build và mở đúng EXE mới', 'Hãy thực sự chạy build_exe.ps1 trên Windows để tạo dist/AI_ExamGuard_Portable.exe hoặc tên có phiên bản nếu tệp cũ đang mở, kiểm tra tệp đầu ra tồn tại và thời gian build, mở đúng tệp mới để thử khởi động, đối chiếu phiên bản trên tiêu đề với /api/runtime và gửi cho tôi đường dẫn đầy đủ cùng dung lượng, không chỉ cung cấp lệnh build mà chưa thực hiện.'),
        ('Thử chạy trên máy Windows khác', 'Hãy chuẩn bị và hướng dẫn tôi kiểm tra EXE trên máy Windows 10 hoặc 11 64 bit khác không cài Node.js hoặc Python nhưng có WebView2, dùng một máy làm giám thị và một máy làm thí sinh trong cùng LAN, kiểm tra IP và cổng thực tế, webcam, nhận diện, cảnh báo đúng máy và học sinh, thử không có Internet nhưng còn LAN với mô hình cục bộ, đồng thời để tôi tự cho phép Firewall trên mạng tin cậy và không đánh dấu thử máy khác là đạt nếu chưa thực hiện.'),
        ('Bàn giao đầy đủ dự án', 'Hãy bàn giao mã nguồn, tệp khóa phụ thuộc, requirements.txt, script build, tệp spec, EXE, hướng dẫn tiếng Việt cho giáo viên và học sinh cùng báo cáo test, giải thích cách lấy địa chỉ máy chủ, thử một máy bằng hai cửa sổ và xử lý lỗi quay mặt không báo, nêu rõ giới hạn không nhận diện điện thoại trực tiếp và không tự kết luận gian lận, rồi tạo gói ZIP phát hành chỉ gồm tài liệu và tệp cần gửi mà không chứa dữ liệu học sinh, khóa bí mật, thư mục phát triển hoặc log cá nhân.'),
    ]),
]


def font_style(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Calibri')


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.7)
sec.left_margin, sec.right_margin = Inches(0.85), Inches(0.85)
sec.footer_distance = Inches(0.32)
font_style(doc.styles['Normal'], 11)
normal_pf = doc.styles['Normal'].paragraph_format
normal_pf.line_spacing = 1.12
normal_pf.space_after = Pt(9)
normal_pf.widow_control = True
font_style(doc.styles['Title'], 23, True)
font_style(doc.styles['Subtitle'], 12)
font_style(doc.styles['Heading 1'], 17, True)
font_style(doc.styles['Heading 2'], 12, True)
for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    pf = doc.styles[name].paragraph_format
    pf.keep_with_next = True
    pf.space_after = Pt(7)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)

doc.core_properties.title = '32 câu lệnh tạo AI ExamGuard và build EXE'
doc.core_properties.subject = 'Hướng dẫn từng bước giao việc cho Codex'
doc.core_properties.author = ''
doc.core_properties.keywords = 'AI ExamGuard, Codex, prompt, Windows, EXE'

doc.add_paragraph('32 câu lệnh tạo AI ExamGuard và build EXE', 'Title')
doc.add_paragraph('Hướng dẫn giao việc từng bước cho Codex', 'Subtitle')
doc.add_paragraph('Tài liệu hướng dẫn tạo bản LAN của AI ExamGuard với giao diện web và vỏ ứng dụng Python, từ thư mục mới đến EXE để gửi chạy trên máy Windows khác; bản cơ bản không cần Firebase và chỉ cảnh báo nghi vấn để giáo viên xác minh.')
doc.add_paragraph('Cách dùng', 'Heading 2')
doc.add_paragraph('Mở một thư mục dự án mới trong Codex, sao chép nguyên đoạn dưới từng bước vào ô chat và gửi lần lượt từ bước 1 đến bước 32 trong cùng cuộc trò chuyện; đây là câu lệnh bằng ngôn ngữ tự nhiên, không phải lệnh gõ vào PowerShell.')
doc.add_paragraph('Chờ Codex hoàn thành và báo kết quả kiểm tra trước khi gửi bước tiếp theo; khi gặp lỗi, xử lý xong lỗi của bước hiện tại trước, còn các bước cần camera hoặc máy thứ hai thì bạn phải phối hợp thử thực tế.')

step = 0
for page_index, (heading, entries) in enumerate(PAGES):
    if page_index:
        doc.add_page_break()
    doc.add_paragraph(heading, 'Heading 1')
    for title, command in entries:
        step += 1
        doc.add_paragraph(f'Bước {step:02d} {title}', 'Heading 2')
        p = doc.add_paragraph(command)
        p.paragraph_format.keep_together = True

assert step == 32
footer = sec.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('Trang ')
r.font.size = Pt(9)
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
doc.settings.element.append(OxmlElement('w:updateFields'))
doc.settings.element[-1].set(qn('w:val'), 'true')
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
print(f'{step} steps; {sum(len(command.split()) for _, entries in PAGES for _, command in entries)} words in commands')
