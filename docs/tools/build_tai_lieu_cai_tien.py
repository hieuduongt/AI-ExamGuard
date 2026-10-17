"""
Tạo Word: Tài liệu phần cải tiến — ôn vấn đáp giám khảo (THCS, không cần biết lập trình).
Đầu ra: docs/tai-lieu-du-thi/AI_ExamGuard_Tai_lieu_phan_cai_tien_theo_gop_y.docx
Chạy: /tmp/egdoc/bin/python docs/tools/build_tai_lieu_cai_tien.py
"""
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_tai_lieu_hoc_sinh import MUTED, NAVY, BLUE, Writer, font_run

OUTPUT = Path(__file__).resolve().parents[1] / "tai-lieu-du-thi" / "AI_ExamGuard_Tai_lieu_phan_cai_tien_theo_gop_y.docx"


class Doc(Writer):
    def _header_footer(self, section):
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        font_run(header.add_run("AI ExamGuard · Ôn vấn đáp phần cải tiến (THCS)"), 9, MUTED, italic=True)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font_run(footer.add_run("Trang "), 9, MUTED)
        self._field(footer, "PAGE", 9)

    def h1(self, text, new_page=False):
        if new_page:
            self.doc.add_page_break()
        self.doc.add_heading(text, level=1)


def cover(w):
    for _ in range(3):
        w.doc.add_paragraph()
    p = w.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("AI ExamGuard"), 36, NAVY, True)
    p = w.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("TÀI LIỆU ÔN VẤN ĐÁP"), 22, BLUE, True)
    p = w.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("CHỈ PHẦN CẢI TIẾN THEO GÓP Ý GIÁM KHẢO"), 16, NAVY, True)
    p = w.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("Dành cho học sinh THCS · Không cần biết lập trình"), 13, MUTED)
    p = w.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("Hiểu bản chất · Biết file nào · Trả lời được câu hỏi sâu"), 12, MUTED, italic=True)
    p = w.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run("Nhóm: Minh · Ngân · Linh"), 12, MUTED)


def part0(w):
    w.h1("0. Em đọc tài liệu này để làm gì?")
    w.p("Giám khảo đã góp ý **bốn hướng cải tiến**. Nhóm đã sửa sản phẩm theo đúng các ý đó. "
        "Tài liệu này **chỉ nói phần cải tiến** (không ôn lại toàn bộ dự án cũ), để em tập trung vấn đáp.")
    w.callout("nho", [
        "**Mỗi bạn trong nhóm đều phải nắm như nhau** — không chia “bạn này trả lời phần này”. Giám khảo hỏi ai cũng phải nói được.",
        "Em **không cần viết code thuộc lòng**. Em cần hiểu: vấn đề là gì → nhóm làm gì → làm ở đâu → vì sao chọn cách đó → demo/bằng chứng thế nào.",
        "Khi giám khảo hỏi sâu: trả lời bằng **ví dụ đời thường** trước, rồi mới nhắc tên file nếu cần.",
        "Câu “an toàn”: “AI không kết luận gian lận; AI chỉ hỗ trợ; giám thị quyết định.”",
    ])
    w.h2("Cách đọc trước vấn đáp (cả nhóm cùng đọc hết)")
    w.steps([
        "Đọc mục 1 (bản đồ 4 góp ý) — nắm toàn cục khoảng 5 phút.",
        "Đọc mục 2 (hai tầng NHÌN / QUYẾT ĐỊNH) — “xương sống” trả lời mọi câu khó về AI.",
        "Đọc lần lượt mục 3 → 6 → 7 → 8 (từng cải tiến + gỡ âm thanh + đường dự phòng). **Không bỏ mục nào.**",
        "Đọc mục 9 (bản đồ file) + mục 11 (câu hỏi khó) — mỗi bạn luyện nói to **cùng một bộ câu hỏi**.",
        "Chạy thử demo mục 10 trên máy thật ít nhất một lần, để ai cũng biết thao tác.",
    ])
    w.h2("Những gì mỗi bạn phải nói được (không phân công)")
    w.table(["Chủ đề", "Phải hiểu / nói được", "Bằng chứng / thao tác demo"], [
        ["Báo nhầm", "Vì sao có hai nút; đếm theo loại để chỉnh ngưỡng; lưu ở hub.", "Bấm Báo nhầm → thẻ số tăng; CSV có cột Kết luận."],
        ["Thử nghiệm + ngưỡng", "Cúi/ngẩng 6 giây vs quay đầu 3 giây; quy trình đếm đúng/nhầm/sót; chỉnh rồi thử lại.", "Kịch bản + bảng số lần 1 / lần 2; ô Quản trị."],
        ["Đoạn hình + biên bản", "Clip tạo lúc AI cảnh báo (không lúc bấm xác nhận); RAM; mặc định tắt; biên bản chỉ confirmed.", "Play clip / lưới ảnh; nút Biên bản đã xác nhận."],
        ["Buổi thử thầy cô", "Nghe người dùng thật; phiếu; 1–2 sửa theo góp ý.", "Ảnh / phiếu / bảng tiếp thu."],
        ["Bản chất AI", "MediaPipe chỉ nhìn; BehaviorEngine đếm; giám thị quyết định; future-custom là ổ cắm trống.", "Giải thích bằng ví von + nhắc tên file nếu được hỏi."],
    ], [3.2, 7.2, 6.2], size=9)


def part1(w):
    w.h1("1. Bản đồ bốn góp ý và nhóm đáp lại thế nào?", new_page=True)
    w.p("Hãy nhớ bảng này như “đề bài giám khảo”. Mỗi dòng = một ý góp + một việc nhóm đã làm.")
    w.table(["#", "Giám khảo muốn (ý chính)", "Nhóm đáp lại trong sản phẩm", "Bằng chứng khi vấn đáp"], [
        ["1", "AI không đúng 100%. Báo nhầm nguy hiểm hơn bỏ sót. Cần nút **Báo nhầm**, đếm để chỉnh thời gian chờ.", "Hai nút **Xác nhận đúng** / **Báo nhầm**; đếm tổng + theo loại; CSV có cột Kết luận.", "Ảnh màn hình thẻ Báo nhầm; bấm thử một cảnh báo."],
        ["2", "Thử có hệ thống ~15 phút theo kịch bản; đếm đúng/nhầm/sót; chỉnh ngưỡng rồi chạy lại; thử thiếu sáng/kính/lệch camera.", "Tách ngưỡng cúi/ngẩng (mặc định 6 giây); ô chỉnh riêng ở Quản trị; quy trình thử + bảng số liệu.", "Kịch bản + bảng trước/sau; video quay quá trình thử/chỉnh."],
        ["3", "Cảnh báo chỉ có chữ thì muộn. Cần **đoạn hình ngắn** quanh lúc xảy ra + biên bản chỉ lần đã xác nhận.", "Clip ~5 giây (hoặc lưới ảnh dự phòng) khi bật công tắc; nút **Biên bản đã xác nhận**.", "Mở chi tiết → nút Xem đoạn hình; xuất biên bản mẫu."],
        ["4", "Thử buổi thật với thầy cô làm giám thị, hỏi ý kiến để cải UI giám thị.", "Không phải sửa lõi AI; là tổ chức buổi thử + phiếu + 1–2 sửa dễ.", "Ảnh buổi thử; phiếu; danh sách đã sửa theo ý thầy cô."],
    ], [0.7, 4.5, 5.5, 5.9], size=9)
    w.callout("canh", [
        "**Không nói** “có tiếng bíp cảnh báo”. Nhóm đã **gỡ âm thanh** theo định hướng hiện tại. Nói: giám thị thấy màu máy, dải cảnh báo mới, mở chi tiết để kết luận.",
        "Góp ý 2: cách làm là **thử trực tiếp theo kịch bản** trước webcam, đếm số, chỉnh ngưỡng, thử lại, rồi **quay lại quá trình** thử/chỉnh làm bằng chứng.",
    ])


def part2(w):
    w.h1("2. Bản chất cần nhớ trước khi đi sâu từng cải tiến", new_page=True)
    w.h2("2.1. Ba vai: Người nhìn — Trọng tài — Giám thị")
    w.p("Nhiều câu hỏi khó của giám khảo đều xoay quanh chỗ này. Học thuộc **ý**, không cần thuộc code.")
    w.table(["Vai", "Ví von", "Làm gì", "Không làm gì", "Chỗ trong dự án"], [
        ["Tầng NHÌN", "Người gác cổng quan sát", "Nhìn camera → vài con số (số mặt, góc đầu, hướng mắt).", "Không kết luận gian lận; không nhận diện “em A/em B”.", "`web/src/ai/` (MediaPipe, vision-provider, face-metrics)"],
        ["Tầng QUYẾT ĐỊNH CẢNH BÁO", "Trọng tài bấm đồng hồ", "Đếm thời gian; đủ ngưỡng thì tạo cảnh báo gọi giám thị.", "Không tự tuyên bố thí sinh sai phạm.", "`web/src/domain/behavior-engine.ts`, `thresholds.ts`"],
        ["CON NGƯỜI", "Giám thị phòng thi", "Tới tận nơi; bấm Xác nhận đúng / Báo nhầm; xuất biên bản.", "Không để AI thay mình kết luận cuối.", "Trang `/giam-thi`"],
    ], [2.6, 2.8, 4.2, 3.6, 3.4], size=9)
    w.callout("vi", "Giống camera cửa hàng: camera chỉ quay hình; chuông kêu khi có chuyển động lâu; bảo vệ mới quyết định có chuyện thật hay không.")
    w.h2("2.2. Vì sao phải tách hai tầng?")
    w.bullets([
        "Nếu gộp “nhìn” và “kết luận” vào một chỗ: mỗi lần đổi AI phải sửa cả luật cảnh báo → dễ hỏng, khó giải thích.",
        "Tách ra: đổi cách nhìn (MediaPipe → thuật toán khác) thì luật đếm giây / nút Báo nhầm / biên bản **có thể giữ**.",
        "Giám khảo hỏi “AI có kết luận gian lận không?” → trả lời chắc: **Không. AI chỉ hỗ trợ quan sát và cảnh báo.**",
    ])
    w.h2("2.3. File “ổ cắm” tầng nhìn: vision-provider")
    w.p("File **`web/src/ai/vision-provider.ts`** là **cửa duy nhất** từ kết quả mô hình nhìn sang các con số mà BehaviorEngine hiểu. "
        "Trong đó có nhánh dự phòng `future-custom` (ổ cắm trống, cố ý chưa gắn thuật toán mới).")
    w.table(["Câu hỏi giám khảo", "Trả lời ngắn"], [
        ["MediaPipe có kết luận gian lận không?", "Không. Nó chỉ nhìn mặt/đầu/mắt và đưa ra vài con số."],
        ["Em đã làm AI mới chưa?", "Chưa. Nhóm để sẵn chỗ `future-custom` để sau này gắn, không viết lại cả hệ thống."],
        ["Sau này gắn thuật toán mới sửa đâu?", "Viết hàm trả về cùng kiểu con số VisionMetrics, gắn vào nhánh future-custom, đổi nhà cung cấp đang dùng; BehaviorEngine và trang giám thị giữ nếu không thêm loại cảnh báo mới."],
    ], [5.5, 11.1], size=10)


def part3(w):
    w.h1("3. Cải tiến 1 — Nút Báo nhầm và đếm để chỉnh ngưỡng", new_page=True)
    w.h2("3.1. Vấn đề giám khảo chỉ ra")
    w.p("AI sẽ có **báo nhầm** (cúi nhặt bút, ngẩng suy nghĩ) và **bỏ sót** (gian lận ngoài khung hình). "
        "**Báo nhầm nguy hiểm hơn** vì thí sinh bị nghi oan, mất tập trung. "
        "Trước đây chỉ có kiểu “đã xem”, không biết đúng hay nhầm → **không đo được thì không chỉnh ngưỡng được**.")
    w.h2("3.2. Mục tiêu / mục đích")
    w.bullets([
        "**Mục tiêu:** giám thị sau khi kiểm tra tận nơi ghi được “đúng” hoặc “báo nhầm”.",
        "**Mục đích:** đếm báo nhầm (đặc biệt theo từng loại) để biết chỗ nào cần tăng thời gian chờ (ngưỡng giây).",
        "**Không nhằm:** để AI tự học hay tự kết luận — vẫn là người bấm.",
    ])
    w.h2("3.3. Ý tưởng giải pháp (không cần biết lập trình)")
    w.p("Mỗi cảnh báo được thêm một “nhãn kết luận” sau khi giám thị kiểm tra:")
    w.bullets([
        "**Xác nhận đúng** = có nghi vấn thật (tiếng Anh trong dữ liệu: `confirmed`).",
        "**Báo nhầm** = AI báo oan (tiếng Anh: `false_alarm`).",
        "Chưa bấm = chưa kết luận.",
    ])
    w.h2("3.4. Sửa ở những file nào? Vì sao ở đó?")
    w.table(["Đường dẫn file", "Chức năng trong cải tiến này", "Vì sao sửa đây (không sửa chỗ khác)"], [
        ["`web/src/domain/exam-types.ts`", "Khai báo kiểu kết luận `AlertVerdict` và trường `verdict?` trên cảnh báo.", "Đây là “từ điển dữ liệu” chung. Mọi màn hình đọc cùng một kiểu — sửa một lần, tránh mỗi trang hiểu khác nhau."],
        ["`shared/hub-contract.json`", "Ghi danh sách `verdicts` web ↔ Python phải thống nhất.", "Hợp đồng giữa hai phía. Nếu chỉ sửa web mà không sửa hợp đồng, bài kiểm thử và Python sẽ lệch."],
        ["`desktop-hub/examguard_desktop/exam_store.py`", "Hàm lưu `reviewed` + `verdict`; chỉ chấp nhận hai giá trị hợp lệ.", "Hub là nơi giữ dữ liệu ca thi trong RAM. Kết luận phải lưu ở đây thì mọi máy giám thị đọc lại được."],
        ["`desktop-hub/examguard_desktop/http_handler.py`", "API PATCH nhận `verdict`; giá trị lạ → lỗi 400.", "Cửa nhận lệnh từ trang web. Không để chữ lạ (ví dụ “abc”) làm bẩn dữ liệu."],
        ["`web/src/services/exam-sync.ts`", "`markEventReviewed(id, true, verdict)` gửi PATCH kèm kết luận.", "Một cửa gửi đi — trang giám thị không gọi thẳng Python, tránh lặp code."],
        ["`web/src/services/local-demo/local-store.ts`", "Lưu verdict khi chạy thử không qua hub.", "Chế độ demo trong trình duyệt cũng phải có cùng hành vi."],
        ["`web/src/services/firebase/firestore-sync.ts`", "Cập nhật verdict nếu dùng Firebase.", "Đủ ba kênh (local / LAN / cloud) cùng luật."],
        ["`web/src/features/proctor/components/event-detail-drawer.tsx`", "Hai nút Xác nhận đúng / Báo nhầm; chỗ xem clip/ảnh.", "Đây là nơi giám thị thao tác sau khi đọc chi tiết — đúng chỗ “sau khi kiểm tra” mà giám khảo yêu cầu."],
        ["`web/src/features/proctor/components/stat-cards.tsx`", "Thẻ số Báo nhầm + dòng “Báo nhầm theo loại”.", "Đếm để chỉnh ngưỡng phải **nhìn thấy ngay**, không bắt mở Excel mới biết."],
        ["`web/src/features/proctor/event-filters.ts`", "Hàm đếm false_alarm / theo loại.", "Tách logic đếm khỏi giao diện — dễ kiểm thử, ít lỗi."],
        ["`web/src/features/proctor/csv-export.ts`", "Cột Kết luận; hàm biên bản chỉ `confirmed`.", "Xuất báo cáo / biên bản dùng chung một cách ghi chữ tiếng Việt."],
        ["`web/src/features/proctor/proctor-screen.tsx`", "Nối nút → markEventReviewed; nút xuất biên bản.", "Màn hình tổng: không nhét logic sâu vào đây, chỉ “bấm nút gọi việc đã có”."],
        ["Các file test: `web/tests/unit/proctor.test.cjs`, `contract.test.cjs`, `desktop-hub/tests/test_hub_api.py`, `test_contract.py`", "Khóa chặt: verdict lạ bị từ chối; CSV có cột Kết luận; hợp đồng khớp.", "Tránh cải tiến hôm nay làm hỏng ngầm ngày mai."],
    ], [4.4, 4.8, 7.4], size=8)
    w.h2("3.5. Vì sao không chỉ sửa giao diện?")
    w.bullets([
        "Nếu chỉ thêm nút trên màn hình mà **không lưu** `verdict` vào hub → tắt trang / máy khác mở lại là mất kết luận.",
        "Nếu chỉ lưu trên một máy trình duyệt → không phải hệ thống phòng thi LAN.",
        "Phải sửa **từ điển dữ liệu → hợp đồng → Python → dịch vụ gửi → giao diện → đếm → CSV** theo chuỗi.",
    ])
    w.h2("3.6. Demo 30 giây khi vấn đáp")
    w.steps([
        "Mở `/giam-thi` → chọn một cảnh báo (hoặc “Tạo cảnh báo thử”).",
        "Bấm **Báo nhầm** → thẻ “Báo nhầm” tăng; có thể thấy theo loại.",
        "Xuất báo cáo CSV → cột **Kết luận** = “Báo nhầm”.",
        "Nói: “Số này dùng để biết loại nào hay báo oan, rồi tăng ngưỡng giây ở trang Quản trị.”",
    ])


def part4(w):
    w.h1("4. Cải tiến 2 — Thử có hệ thống + tách ngưỡng cúi/ngẩng", new_page=True)
    w.h2("4.1. Giám khảo muốn gì? (đọc đúng chữ)")
    w.p("Quay / tổ chức vài đoạn khoảng **15 phút** một bạn làm bài theo **kịch bản**: việc bình thường "
        "(cúi viết nháp, nhìn bàn phím, uống nước) và việc cần báo (quay sang bên, rời chỗ, người đứng sau). "
        "Đếm **đúng / báo nhầm / bỏ sót**. Mỗi lần chỉnh ngưỡng thì **chạy lại đúng các đoạn đó**. "
        "Thử thêm thiếu sáng, đeo kính, ngồi lệch camera.")
    w.callout("nho",
              "Cách làm đúng ý giám khảo: viết kịch bản → **thử trực tiếp** trước webcam → đếm → chỉnh ngưỡng → **thử lại** → "
              "**quay lại quá trình** thử/chỉnh làm bằng chứng (có bảng số liệu).")
    w.h2("4.2. Vì sao tách ngưỡng cúi/ngẩng?")
    w.p("Trước đây **cúi/ngẩng và quay đầu dùng chung một mốc thời gian** (khoảng 3 giây). "
        "Trong giờ thi, cúi viết nháp / nhặt bút / ngẩng suy nghĩ thường lâu hơn 3 giây → **dễ báo nhầm**. "
        "Quay hẳn sang bạn bên cạnh đáng chú ý sớm hơn.")
    w.bullets([
        "**Mục tiêu:** giảm báo nhầm việc bình thường mà vẫn bắt được quay đầu rõ.",
        "**Cách làm:** thêm ngưỡng riêng `headPoseSeconds` (mặc định **6 giây**); quay đầu giữ **3 giây**.",
        "**Chỗ chỉnh:** trang `/quan-tri` → ô “Cúi/ngẩng đầu”.",
    ])
    w.h2("4.3. File nào liên quan ngưỡng? Vì sao?")
    w.table(["Đường dẫn", "Việc làm", "Vì sao ở đây"], [
        ["`web/src/domain/exam-types.ts`", "Thêm `headPoseSeconds` vào bộ ngưỡng.", "Cùng chỗ khai báo mọi ngưỡng — một “bảng luật”."],
        ["`web/src/domain/thresholds.ts`", "Mặc định cúi/ngẩng = 6; hàm đọc/cất ngưỡng đã lưu.", "Số mặc định và cách đọc cấu hình nằm một chỗ."],
        ["`web/src/domain/behavior-engine.ts`", "Loại `head_pose` dùng `headPoseSeconds` (không còn dùng nhầm `headTurnSeconds`).", "Đây mới là chỗ **đếm giây**. Sửa UI mà không sửa đây = nút giả."],
        ["`web/src/features/admin/components/threshold-settings-card.tsx`", "Ô chỉnh riêng trên Quản trị.", "Người vận hành chỉnh được không cần mở code."],
        ["`web/tests/unit/domain.test.cjs`", "Khóa mặc định mới + hành vi đếm.", "Tránh ai đó đổi nhầm về 3 giây mà không biết."],
    ], [4.6, 5.2, 6.8], size=9)
    w.h2("4.4. Vì sao không “học máy” tự chỉnh ngưỡng?")
    w.bullets([
        "Dự án THCS cần **giải thích được** và **kiểm soát được**.",
        "Giám khảo muốn thấy nhóm **đo → chỉnh → đo lại**, không phải hộp đen.",
        "Người (giám thị / quản trị) vẫn là người đặt ngưỡng sau khi có số liệu.",
    ])
    w.h2("4.5. Quy trình thử nghiệm (cả nhóm phải nắm)")
    w.steps([
        "Viết kịch bản có mốc thời gian (mẫu trong kế hoạch cải tiến).",
        "Buổi A ~15 phút ánh sáng bình thường; B thiếu sáng; C kính / ngồi lệch.",
        "Thử lần 1 (webcam) → xuất CSV → điền bảng đúng/nhầm/sót; giám thị bấm Báo nhầm khi oan.",
        "Chỉnh ngưỡng (ví dụ tăng cúi/ngẩng) → **thử lại đúng danh sách tình huống**.",
        "Quay màn hình / góc phòng lúc thử và lúc chỉnh — đây là video bằng chứng.",
        "Viết 3–4 câu kết luận trung thực (kể cả nếu bỏ sót tăng).",
    ])


def part5(w):
    w.h1("5. Cải tiến 3 — Đoạn hình ngắn + biên bản đã xác nhận", new_page=True)
    w.h2("5.1. Vấn đề")
    w.p("Cảnh báo chỉ có chữ và thời điểm → khi giám thị tới nơi, việc thường **đã xong**. "
        "Cần xem lại vài giây trước và trong lúc xảy ra. Cuối buổi cần bản tổng kết **chỉ những lần người đã xác nhận**.")
    w.h2("5.2. Clip được tạo khi nào? Lưu ở đâu?")
    w.callout("canh", [
        "Clip **không** được tạo lúc bấm “Xác nhận đúng”.",
        "Clip gắn vào cảnh báo **ngay khi AI phát cảnh báo** (nếu công tắc lưu đang bật).",
        "Nút xác nhận chỉ gắn nhãn đúng/nhầm lên cảnh báo đã có.",
    ])
    w.bullets([
        "**Nơi lưu:** trong **bộ nhớ RAM** của máy giám thị (hub), nằm trong dữ liệu cảnh báo (`clip` hoặc `snapshots`).",
        "**Không** tự lưu file vào ổ đĩa kiểu thư mục `Videos/`.",
        "**Đóng EXE / tắt hub → mất.** Đó là chủ đích quyền riêng tư.",
        "**Biên bản CSV** chỉ là chữ — **không nhét video vào Excel**.",
    ])
    w.h2("5.3. Cách dùng (học thuộc thao tác)")
    w.steps([
        "Mở `/quan-tri` → bật **“Lưu đoạn hình ngắn khi cảnh báo”** → **Lưu thay đổi** (đẩy lên hub).",
        "Trên máy thí sinh: **bắt đầu lại phiên** (công tắc được đọc lúc bắt đầu).",
        "Khi có cảnh báo AI → đợi khoảng 2 giây (để ghi phần “trong lúc xảy ra”).",
        "`/giam-thi` → mở chi tiết → nút **“Xem lại đoạn hình ngắn”** (hoặc lưới ảnh dự phòng).",
        "Tới nơi kiểm tra → bấm Xác nhận đúng / Báo nhầm.",
        "Cuối buổi → **Biên bản đã xác nhận** (chỉ các lần `confirmed`).",
    ])
    w.h2("5.4. File nào? Vì sao?")
    w.table(["Đường dẫn", "Việc làm", "Vì sao ở đây / không ở chỗ khác"], [
        ["`web/src/features/candidate/hooks/alert-clip-recorder.ts`", "Giữ vòng đệm ~3 giây trước; khi cảnh báo ghi thêm ~2 giây; ghép thành video ngắn.", "Tách việc ghi hình ra file riêng — `use-monitoring-session` đã rất dài; dễ bảo trì và giải thích."],
        ["`web/src/features/candidate/hooks/use-monitoring-session.ts`", "Bật recorder khi công tắc bật; khi emit cảnh báo AI thì chốt clip; gửi kèm sự kiện.", "Đây là “bộ não phiên thí sinh” — chỉ chỗ này thấy webcam thật để ghi."],
        ["`web/src/domain/exam-types.ts`", "Trường `clip?`, `snapshots?` trên cảnh báo.", "Cùng từ điển dữ liệu với verdict."],
        ["`web/src/features/admin/components/privacy-settings-card.tsx`", "Công tắc lưu đoạn hình (mặc định tắt).", "Quyền riêng tư phải do nhà trường bật có chủ đích."],
        ["`web/src/features/admin/admin-screen.tsx` + `web/src/services/hub/hub-settings.ts`", "Lưu local + **đẩy công tắc lên hub** `/api/settings`.", "Trước đây chỉ lưu trên máy mở Quản trị → máy thí sinh khác không biết. Đẩy lên hub để cả phòng dùng chung."],
        ["`desktop-hub/.../exam_store.py` + `http_handler.py`", "GET/POST `/api/settings`; lưu sự kiện kèm `clip`.", "Hub là trung tâm LAN; settings + events phải ở đây."],
        ["`web/src/services/exam-sync.ts`", "Cảnh báo có clip gửi **thẳng** hub (không nhét localStorage).", "Clip lớn dễ vượt hạn mức localStorage của hộp thư đi → mất media mà vẫn tưởng đã gửi."],
        ["`web/src/features/proctor/components/event-detail-drawer.tsx`", "Nút xem clip / lưới ảnh; thông báo nếu không có media.", "Giám thị xem lại đúng chỗ đang kết luận."],
        ["`web/src/features/proctor/components/event-feed.tsx`", "Badge “Có đoạn hình” / “Có ảnh”.", "Nhìn danh sách biết cảnh báo nào xem lại được."],
        ["`web/src/features/proctor/csv-export.ts`", "`downloadConfirmedCsv` lọc `verdict === 'confirmed'`.", "Biên bản sạch — đúng yêu cầu “chỉ lần đã xác nhận”."],
        ["`shared/hub-contract.json` + `config.py`", "`maxBodyBytes` ~1,5 MB.", "Clip base64 cần chỗ gửi lớn hơn 1 MB cũ."],
    ], [4.5, 4.8, 7.3], size=8)
    w.h2("5.5. Câu trả lời khó về quyền riêng tư")
    w.bullets([
        "Mặc định **tắt** — không ghi hình nếu nhà trường không bật.",
        "Chỉ vài giây quanh cảnh báo, độ phân giải thấp; chỉ trong RAM máy giám thị.",
        "Thí sinh vẫn phải đồng ý giám sát trước khi bật camera.",
        "Không dùng để nhận diện danh tính.",
    ])


def part6(w):
    w.h1("6. Cải tiến 4 — Buổi thử thật với thầy cô (không phải code)", new_page=True)
    w.p("Giám khảo muốn nhóm **nghe người dùng thật**. Đây là phần tổ chức và tiếp thu góp ý — **không phải thêm file AI**, nhưng **mọi bạn đều phải kể được** đã thử gì và học được gì.")
    w.steps([
        "Mời một thầy cô làm giám thị 30–45 phút; 2–3 bạn làm thí sinh.",
        "Phát phiếu hỏi ý kiến (mẫu trong kế hoạch cải tiến).",
        "Ghi lại 1–2 ý dễ sửa (chữ nút, bố cục) → sửa trên sản phẩm nếu kịp.",
        "Trong hồ sơ: bảng “Góp ý → Đã làm → Bằng chứng → Còn lại”.",
    ])
    w.callout("nho", "Khi vấn đáp: đừng nói “em chưa kịp thử” hoặc “phần đó bạn kia phụ trách”. Nói thật những gì nhóm đã thử + bài học từ phiếu.")


def part7(w):
    w.h1("7. Việc đã gỡ: cảnh báo bằng âm thanh", new_page=True)
    w.p("Theo định hướng hiện tại, nhóm **không dùng** tiếng bíp.")
    w.bullets([
        "Đã xóa `web/src/features/proctor/hooks/use-alert-sound.ts` và file âm thanh trong `web/public/sounds/`.",
        "Đã gỡ nút bật/tắt âm thanh khỏi tiêu đề trang giám thị.",
        "Vẫn còn **dải cảnh báo mới** (nhìn thấy, không kêu) và đổi màu máy.",
    ])
    w.callout("canh", "Nếu giám khảo hỏi “sao không có chuông?” → “Nhóm ưu tiên giao diện nhìn thấy, tránh ồn phòng thi khi trình diễn; giám thị vẫn mở chi tiết để kết luận đúng/nhầm.”")


def part8(w):
    w.h1("8. Đường dự phòng thuật toán nhìn (vision-provider) — trả lời câu hỏi sâu", new_page=True)
    w.p("Đây là phần giám khảo thích hỏi để xem em có hiểu **giới hạn AI** và **thiết kế mở rộng** không.")
    w.h2("8.1. Vì sao cần?")
    w.bullets([
        "MediaPipe chưa hoàn hảo (phòng tối, kính, ngồi lệch).",
        "Sau này có thể muốn thuật toán nhìn khác — không muốn đập đi xây lại cả BehaviorEngine, nút Báo nhầm, biên bản.",
        "Thể hiện tư duy: biết chỗ yếu và để **ổ cắm** có chủ đích.",
    ])
    w.h2("8.2. File trung tâm")
    w.p("**`web/src/ai/vision-provider.ts`**")
    w.bullets([
        "Đang bật: `mediapipe-face-landmarker`.",
        "Chỗ trống: `future-custom` — **chưa gắn thuật toán thật**; bật nhầm sẽ báo lỗi rõ (an toàn).",
        "Vòng lặp AI ở `use-monitoring-session.ts` gọi `metricsFromVisionResult(...)` — **không** gọi thẳng MediaPipe.",
    ])
    w.h2("8.3. Vì sao không sửa BehaviorEngine để “đổi AI”?")
    w.p("BehaviorEngine chỉ hiểu **vài con số + thời gian**. Nó không cần biết phía sau là MediaPipe hay mô hình khác. "
        "Đổi AI ở tầng nhìn đúng chỗ; giữ luật cảnh báo ổn định.")
    w.callout("vi", "Ổ USB trên máy tính: hôm nay cắm chuột, ngày mai cắm bàn phím khác — không phải thay cả bo mạch.")


def part9(w):
    w.h1("9. Bản đồ file cải tiến (học thuộc cột “nói một câu”)", new_page=True)
    w.p("Mỗi bạn nên đọc hết bảng này. Không cần thuộc từng dòng code — chỉ cần **nói được một câu** về từng file nếu giám khảo hỏi “sửa ở đâu?”.")
    w.table(["Đường dẫn", "Một câu khi vấn đáp"], [
        ["`web/src/domain/exam-types.ts`", "“Em khai báo thêm kết luận giám thị, clip/ảnh, ngưỡng cúi/ngẩng.”"],
        ["`shared/hub-contract.json`", "“Hợp đồng web–Python: verdicts, giới hạn dung lượng gửi.”"],
        ["`desktop-hub/.../exam_store.py`", "“Hub lưu cảnh báo + kết luận + cài đặt lưu đoạn hình trong RAM.”"],
        ["`desktop-hub/.../http_handler.py`", "“API nhận PATCH verdict và GET/POST settings.”"],
        ["`web/src/domain/behavior-engine.ts`", "“Chỗ đếm giây; cúi/ngẩng dùng ngưỡng 6 giây riêng.”"],
        ["`web/src/domain/thresholds.ts`", "“Số mặc định và ô cấu hình ngưỡng.”"],
        ["`web/src/ai/vision-provider.ts`", "“Cửa tầng nhìn + ổ cắm future-custom.”"],
        ["`web/src/features/candidate/hooks/alert-clip-recorder.ts`", "“Ghi ~3 giây trước + ~2 giây trong lúc cảnh báo.”"],
        ["`web/src/features/candidate/hooks/use-monitoring-session.ts`", "“Bộ não phiên thí sinh: AI, clip, gửi cảnh báo.”"],
        ["`web/src/features/proctor/components/event-detail-drawer.tsx`", "“Hai nút kết luận + nút xem đoạn hình.”"],
        ["`web/src/features/proctor/components/stat-cards.tsx`", "“Đếm báo nhầm theo loại để chỉnh ngưỡng.”"],
        ["`web/src/features/proctor/csv-export.ts`", "“Biên bản chỉ lần xác nhận đúng.”"],
        ["`web/src/features/admin/.../privacy-settings-card.tsx`", "“Công tắc lưu đoạn hình, mặc định tắt.”"],
        ["`web/src/services/hub/hub-settings.ts`", "“Đẩy/đọc công tắc lưu hình qua hub cho cả phòng.”"],
        ["`web/src/services/exam-sync.ts`", "“Gửi cảnh báo; clip gửi thẳng tránh mất trong hộp thư đi.”"],
    ], [5.5, 11.1], size=9)


def part10(w):
    w.h1("10. Kịch bản demo vấn đáp (5–7 phút)", new_page=True)
    w.p("Ai đứng nói cũng được — **nội dung dưới đây cả nhóm phải nói được**. Có thể luân phiên demo trên máy, nhưng không được bảo “phần này em không biết”.")
    w.steps([
        "Mở đầu (~1 phút): tóm 4 góp ý và nhóm đã đáp lại thế nào (xem mục 1).",
        "Demo Báo nhầm: bấm nút → thẻ số tăng → CSV có cột Kết luận; nói vì sao đếm theo loại.",
        "Demo đoạn hình (nếu đã bật công tắc): badge “Có đoạn hình” → Play → giải thích RAM, mặc định tắt, không quay lúc bấm xác nhận.",
        "Giải thích ngưỡng: cúi/ngẩng 6 giây vs quay đầu 3 giây; chỉ bảng số lần 1 → chỉnh → lần 2.",
        "Nói quyền riêng tư + buổi thử thầy cô + 1 ý đã sửa theo góp ý (nếu có).",
        "Chốt sẵn 3 câu: “AI không kết luận gian lận.” / “future-custom là ổ cắm trống.” / “Clip gắn lúc AI cảnh báo, không lúc bấm xác nhận.”",
    ])


def part11(w):
    w.h1("11. Ngân hàng câu hỏi khó (học trả lời)", new_page=True)
    items = [
        ("Vì sao thêm Báo nhầm chứ không chỉ “Đã xem”?",
         "Vì “đã xem” không cho biết đúng hay oan. Phải đếm báo nhầm theo loại mới biết chỗ nào tăng ngưỡng thời gian."),
        ("Báo nhầm nguy hiểm hơn bỏ sót vì sao?",
         "Thí sinh bị nghi oan, mất tập trung, ảnh hưởng tâm lý và công bằng. Bỏ sót xấu nhưng giám thị vẫn có thể quan sát thêm; nghi oan thì hại trực tiếp thí sinh."),
        ("MediaPipe có kết luận gian lận không?",
         "Không. MediaPipe chỉ nhìn và đưa ra vài con số. BehaviorEngine đếm giờ tạo cảnh báo. Giám thị quyết định."),
        ("Đường dự phòng future-custom là gì? Em đã làm AI mới chưa?",
         "Là chỗ dành sẵn để gắn thuật toán nhìn khác sau này, không viết lại cả hệ thống. Chưa gắn thuật toán mới — cố ý để trống an toàn. Hiện vẫn MediaPipe."),
        ("Giám khảo bảo quay đoạn 15 phút — em hiểu thế nào?",
         "Viết kịch bản, thử trực tiếp trước webcam, đếm đúng/nhầm/sót, chỉnh ngưỡng rồi thử lại, quay lại quá trình thử/chỉnh làm bằng chứng."),
        ("Clip lưu ở đâu? Xác nhận đúng có tạo clip không?",
         "Clip tạo khi AI cảnh báo (nếu bật công tắc), nằm RAM máy giám thị trong dữ liệu cảnh báo. Bấm xác nhận chỉ gắn nhãn đúng/nhầm, không quay thêm."),
        ("Vì sao cúi/ngẩng 6 giây còn quay đầu 3 giây?",
         "Cúi viết nháp là việc bình thường, cần thêm thời gian để giảm báo oan; quay hẳn sang bên đáng chú ý sớm hơn."),
        ("Vì sao không lưu video cả buổi thi?",
         "Quyền riêng tư và dung lượng. Chỉ vài giây quanh cảnh báo, mặc định tắt, chỉ khi nhà trường bật."),
        ("Biên bản gồm những gì?",
         "Chỉ cảnh báo giám thị đã bấm Xác nhận đúng. Không gồm báo nhầm và chưa kiểm tra."),
        ("Vì sao bỏ tiếng bíp?",
         "Ưu tiên tín hiệu nhìn thấy, tránh ồn khi trình diễn/thử; vẫn còn dải cảnh báo mới và màu máy."),
        ("Em làm phần nào trong nhóm?",
         "Có thể nói mình góp sức chỗ nào (code / thử số / buổi thầy cô), nhưng **vẫn phải giải thích được cả bốn cải tiến**. Không được bảo “phần đó bạn kia phụ trách nên em không biết”."),
        ("Nếu chỉnh ngưỡng xong báo nhầm giảm nhưng bỏ sót tăng?",
         "Ghi đúng số và giải thích đánh đổi. Giám khảo thích sự trung thực hơn bảng số “đẹp giả”."),
    ]
    for i, (q, a) in enumerate(items, 1):
        w.qa(i, q, a)


def part12(w):
    w.h1("12. Checklist nộp / trước giờ vấn đáp", new_page=True)
    w.bullets([
        "☐ Mỗi bạn nói được **cả bốn cải tiến** trong khoảng 2 phút, không nhìn giấy.",
        "☐ Mỗi bạn tự demo được: Báo nhầm → số tăng; (nếu có) Play đoạn hình; Biên bản đã xác nhận.",
        "☐ Cả nhóm cùng nắm bảng số lần 1 / lần 2 (dù chưa đẹp).",
        "☐ Cả nhóm cùng nắm ảnh/phiếu buổi thử hoặc nói rõ đã làm gì với thầy cô.",
        "☐ Không nhắc âm thanh cảnh báo.",
        "☐ Thuộc câu: MediaPipe nhìn — BehaviorEngine đếm — Giám thị quyết định — vision-provider là ổ cắm.",
        "☐ Không dùng câu “phần đó bạn kia phụ trách”.",
    ])
    w.callout("nho", "Ba câu chốt: **Làm đúng 4 góp ý. Mỗi góp ý có bằng chứng. AI không kết luận gian lận.**")


def build():
    w = Doc()
    cover(w)
    part0(w)
    part1(w)
    part2(w)
    part3(w)
    part4(w)
    part5(w)
    part6(w)
    part7(w)
    part8(w)
    part9(w)
    part10(w)
    part11(w)
    part12(w)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    w.doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
