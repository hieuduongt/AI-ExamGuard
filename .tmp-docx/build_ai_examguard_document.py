from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "tai-lieu" / "AI_ExamGuard_Mo_ta_tinh_nang_vai_tro_va_cong_nghe.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5C6B7A"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F2F7FC"
PALE_GREEN = "EDF8F2"
GREEN = "16794A"
GOLD = "A66A00"
BODY = "1F2937"
WHITE = "FFFFFF"
BORDER = "D5DDE8"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_table(table, widths_dxa, header_fill=LIGHT_BLUE):
    assert sum(widths_dxa) == 9360
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(widths_dxa[col_index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[col_index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            if row_index == 0:
                shade_cell(cell, header_fill)
        if row_index == 0:
            set_repeat_table_header(row)


def add_numbering_definition(document, *, kind, marker):
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), kind)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), marker)
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "271")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    if kind == "bullet":
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Arial")
        fonts.set(qn("w:hAnsi"), "Arial")
        r_pr.append(fonts)
        lvl.append(r_pr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def assign_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_bullet(document, num_id, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    assign_numbering(paragraph, num_id)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        first.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_numbered(document, num_id, title, detail):
    paragraph = document.add_paragraph()
    assign_numbering(paragraph, num_id)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    title_run = paragraph.add_run(f"{title}: ")
    title_run.bold = True
    paragraph.add_run(detail)
    return paragraph


def set_paragraph_shading(paragraph, fill, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if left_border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        p_bdr.append(left)


def add_callout(document, label, text, fill=PALE_BLUE, accent=BLUE):
    paragraph = document.add_paragraph(style="Callout")
    set_paragraph_shading(paragraph, fill, accent)
    lead = paragraph.add_run(f"{label}: ")
    lead.bold = True
    lead.font.color.rgb = RGBColor.from_string(accent)
    paragraph.add_run(text)
    return paragraph


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_run_font(run, size=9, color=MUTED)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string("2B5163")
    subtitle.font.italic = False
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    lead = styles.add_style("Lead", 1)
    lead.base_style = normal
    lead.font.size = Pt(12.5)
    lead.font.color.rgb = RGBColor.from_string(NAVY)
    lead.paragraph_format.space_after = Pt(10)
    lead.paragraph_format.line_spacing = 1.25

    callout = styles.add_style("Callout", 1)
    callout.base_style = normal
    callout.font.size = Pt(10.5)
    callout.paragraph_format.left_indent = Inches(0.14)
    callout.paragraph_format.right_indent = Inches(0.08)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(10)
    callout.paragraph_format.line_spacing = 1.2


def configure_section_geometry(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_page(document):
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    configure_section_geometry(section)
    section.different_first_page_header_footer = False
    section.header.paragraphs[0].text = ""
    section.even_page_header.paragraphs[0].text = ""
    section.footer.paragraphs[0].text = ""
    section.even_page_footer.paragraphs[0].text = ""


def populate_running_header(header):
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("AI EXAMGUARD  |  TÀI LIỆU MÔ TẢ HỆ THỐNG")
    set_run_font(run, size=8.5, color=MUTED, bold=True)


def populate_running_footer(footer):
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("AI ExamGuard  •  Trang ")
    set_run_font(run, size=9, color=MUTED)
    add_page_field(p)


def start_body_section(document):
    section = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section_geometry(section)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.even_page_header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.even_page_footer.is_linked_to_previous = False
    populate_running_header(section.header)
    populate_running_header(section.even_page_header)
    populate_running_footer(section.footer)
    populate_running_footer(section.even_page_footer)


def add_cover(document):
    for _ in range(3):
        document.add_paragraph().paragraph_format.space_after = Pt(18)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    set_run_font(kicker.add_run("BÁO CÁO MÔ TẢ HỆ THỐNG"), size=11, color=GOLD, bold=True)
    document.add_paragraph("AI EXAMGUARD", style="Title")
    document.add_paragraph("Mô tả tính năng, vai trò và công nghệ sử dụng", style="Subtitle")
    tagline = document.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline.paragraph_format.space_before = Pt(8)
    tagline.paragraph_format.space_after = Pt(54)
    set_run_font(tagline.add_run("Giám sát thi thông minh • Xử lý AI tại thiết bị • Con người quyết định"), size=10.5, color=MUTED, italic=True)

    info = document.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.space_after = Pt(6)
    set_run_font(info.add_run("Phiên bản tài liệu 1.0"), size=10.5, color=NAVY, bold=True)
    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_after = Pt(4)
    set_run_font(date.add_run("Ngày 01 tháng 09 năm 2026"), size=10, color=MUTED)
    url = document.add_paragraph()
    url.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(url.add_run("https://ai-examguard.vercel.app"), size=10, color=BLUE)


def add_intro(document):
    document.add_heading("1. Tổng quan hệ thống", level=1)
    document.add_paragraph(
        "AI ExamGuard là hệ thống web hỗ trợ giám thị phát hiện sớm các tín hiệu bất thường trong phòng thi. "
        "Ứng dụng sử dụng webcam và mô hình thị giác máy tính chạy trực tiếp trên máy thí sinh để phân tích khuôn mặt, "
        "hướng nhìn và tư thế đầu; hệ thống chỉ chuyển trạng thái cùng sự kiện cảnh báo tới màn hình giám thị."
    ).style = "Lead"
    add_callout(
        document,
        "Nguyên tắc cốt lõi",
        "AI chỉ tạo tín hiệu hỗ trợ quan sát, không tự kết luận thí sinh vi phạm. Giám thị là người xem xét bối cảnh và đưa ra quyết định cuối cùng theo quy chế thi.",
    )
    document.add_heading("1.1. Mục tiêu", level=2)
    document.add_paragraph(
        "Hệ thống hướng tới việc giảm tải thao tác theo dõi thủ công, giúp giám thị ưu tiên đúng vị trí cần kiểm tra, "
        "đồng thời bảo vệ quyền riêng tư bằng cách xử lý video tại chỗ và hạn chế truyền dữ liệu không cần thiết."
    )
    document.add_heading("1.2. Phạm vi sử dụng", level=2)
    add_bullet(document, BULLET_ID, "Dùng trong phòng thi có mỗi thí sinh một máy tính và một webcam tích hợp hoặc webcam rời.")
    add_bullet(document, BULLET_ID, "Hỗ trợ ba khu vực nghiệp vụ: thí sinh, giám thị và quản trị hệ thống.")
    add_bullet(document, BULLET_ID, "Hoạt động trên trình duyệt hiện đại qua kết nối HTTPS; phiên bản hiện tại được triển khai trên Vercel.")


def add_features(document):
    document.add_heading("2. Các tính năng chính", level=1)
    document.add_heading("2.1. Khu vực thí sinh", level=2)
    add_bullet(document, BULLET_ID, "Khai báo phiên thi: nhập họ tên, mã phòng và số máy/vị trí trước khi bắt đầu.", "Khai báo phiên thi:")
    add_bullet(document, BULLET_ID, "Xác nhận thông báo giám sát: chỉ cho phép bắt đầu khi thí sinh đã đọc và đồng ý bật webcam.", "Xác nhận thông báo giám sát:")
    add_bullet(document, BULLET_ID, "Kiểm tra webcam: xin quyền camera, hiển thị hình ảnh cục bộ và báo rõ nguyên nhân khi camera bị chặn, không tồn tại hoặc đang được ứng dụng khác sử dụng.", "Kiểm tra webcam:")
    add_bullet(document, BULLET_ID, "Phân tích thời gian thực: hiển thị số khuôn mặt, góc quay đầu, mức cúi/ngẩng, độ lệch hướng mắt và tốc độ xử lý AI.", "Phân tích thời gian thực:")
    add_bullet(document, BULLET_ID, "Phản hồi tại chỗ: thể hiện trạng thái bình thường, cần theo dõi hoặc ưu tiên kiểm tra bằng màu xanh, vàng và đỏ.", "Phản hồi tại chỗ:")

    document.add_heading("2.2. Nhóm tín hiệu được theo dõi", level=2)
    signals = [
        ("Không có khuôn mặt", "Phát hiện thí sinh rời khỏi vùng quan sát quá thời gian cho phép."),
        ("Nhiều khuôn mặt", "Phát hiện có từ hai khuôn mặt trở lên xuất hiện trước webcam."),
        ("Quay đầu kéo dài", "Ước tính góc quay sang trái hoặc phải và chỉ cảnh báo khi vượt ngưỡng liên tục."),
        ("Cúi hoặc ngẩng đầu", "Theo dõi tư thế đầu bất thường dựa trên các điểm mốc khuôn mặt."),
        ("Hướng mắt lệch", "Ước tính hướng nhìn rời vùng màn hình trong một khoảng thời gian đủ dài."),
        ("Rời cửa sổ bài thi", "Ghi nhận khi trang thi bị ẩn hoặc thí sinh chuyển sang cửa sổ khác."),
        ("Thao tác bàn phím hạn chế", "Ghi nhận các tổ hợp sao chép, dán, cắt hoặc in như Ctrl/Cmd+C, V, X và P."),
    ]
    for title, detail in signals:
        add_bullet(document, BULLET_ID, f"{title}: {detail}", f"{title}:")

    add_callout(
        document,
        "Giảm cảnh báo nhiễu",
        "Behavior Engine sử dụng bộ đệm thời gian và khoảng nghỉ giữa hai cảnh báo cùng loại. Ngưỡng mặc định gồm: vắng mặt 5 giây, nhiều khuôn mặt 2 giây, quay/cúi/ngẩng 4 giây, hướng mắt lệch 5 giây và khoảng nghỉ 12 giây.",
        fill="FFF8E8",
        accent=GOLD,
    )

    document.add_heading("2.3. Dashboard giám thị", level=2)
    add_bullet(document, BULLET_ID, "Theo dõi sơ đồ trạng thái của toàn bộ máy thi theo ba mức: ổn định, cảnh báo và ưu tiên kiểm tra.")
    add_bullet(document, BULLET_ID, "Nhận luồng sự kiện, tìm kiếm theo máy hoặc thí sinh, lọc theo mức độ và trạng thái chưa xác minh.")
    add_bullet(document, BULLET_ID, "Mở chi tiết cảnh báo để xem thời gian, phòng thi, nguồn phát hiện và nội dung tín hiệu.")
    add_bullet(document, BULLET_ID, "Đánh dấu sự kiện đã xác minh; hệ thống luôn nhắc giám thị quan sát trực tiếp trước khi xử lý.")
    add_bullet(document, BULLET_ID, "Xuất nhật ký cảnh báo thành tệp CSV phục vụ tổng hợp và báo cáo.")

    document.add_heading("2.4. Khu vực quản trị", level=2)
    add_bullet(document, BULLET_ID, "Cấu hình tên phòng thi và số lượng máy thí sinh.")
    add_bullet(document, BULLET_ID, "Điều chỉnh các ngưỡng thời gian của từng loại cảnh báo và khoảng nghỉ chống lặp.")
    add_bullet(document, BULLET_ID, "Thiết lập yêu cầu đồng ý giám sát và tùy chọn lưu ảnh tại thời điểm cảnh báo; lưu ảnh mặc định được tắt.")
    add_bullet(document, BULLET_ID, "Lưu cấu hình cục bộ và xuất cấu hình dưới dạng JSON để sao lưu hoặc chuyển sang máy khác.")
    add_bullet(document, BULLET_ID, "Hiển thị trạng thái kết nối Firebase để người quản trị nhận biết hệ thống đang dùng đồng bộ thời gian thực hay chế độ demo cục bộ.")

    document.add_heading("2.5. Dữ liệu và quyền riêng tư", level=2)
    add_bullet(document, BULLET_ID, "Video được phân tích trực tiếp trong trình duyệt của máy thí sinh và không truyền liên tục lên dashboard.")
    add_bullet(document, BULLET_ID, "Dữ liệu gửi đi chủ yếu gồm trạng thái máy, chỉ số kỹ thuật và sự kiện cảnh báo.")
    add_bullet(document, BULLET_ID, "Việc lưu ảnh cảnh báo chỉ nên bật khi đã có cơ sở pháp lý, thời hạn lưu, phân quyền truy cập và quy trình xóa dữ liệu rõ ràng.")


def add_roles(document):
    document.add_heading("3. Vai trò và trách nhiệm", level=1)
    document.add_paragraph(
        "Mô hình vận hành phân tách rõ người sử dụng, người giám sát, người cấu hình và hệ thống AI. "
        "Sự phân tách này giúp tránh việc AI tự động thay thế quyết định nghiệp vụ."
    )

    table = document.add_table(rows=1, cols=3)
    headers = ["Vai trò", "Trách nhiệm chính", "Quyền và giới hạn"]
    for index, text in enumerate(headers):
        p = table.rows[0].cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=10, color=NAVY, bold=True)
    role_rows = [
        ("Thí sinh", "Khai báo phiên thi, xác nhận thông báo, bật webcam và giữ đúng vị trí trong suốt thời gian thi.", "Chỉ xem trạng thái và sự kiện trên máy của mình; không kết luận hoặc sửa cảnh báo."),
        ("Giám thị", "Theo dõi dashboard, ưu tiên vị trí cần kiểm tra, xem chi tiết và xác minh sự kiện.", "Là người đánh giá bối cảnh và ra quyết định cuối cùng theo quy chế thi."),
        ("Quản trị hệ thống", "Thiết lập phòng thi, ngưỡng cảnh báo, quyền riêng tư, kênh dữ liệu và xuất cấu hình.", "Không được đặt ngưỡng hoặc chính sách lưu dữ liệu khi chưa kiểm thử và phê duyệt."),
        ("Hệ thống AI", "Phân tích webcam tại chỗ, tính chỉ số và tạo cảnh báo khi tín hiệu vượt ngưỡng thời gian.", "Chỉ hỗ trợ quan sát; không nhận diện danh tính và không tự tuyên bố hành vi gian lận."),
    ]
    for role, responsibility, authority in role_rows:
        cells = table.add_row().cells
        for index, text in enumerate((role, responsibility, authority)):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index == 0:
                set_run_font(p.add_run(text), size=10.5, color=DARK_BLUE, bold=True)
            else:
                set_run_font(p.add_run(text), size=10, color=BODY)
    configure_table(table, [1560, 4020, 3780])

    document.add_heading("3.1. Quy trình phối hợp", level=2)
    workflow = [
        ("Chuẩn bị", "Quản trị viên cấu hình phòng thi, ngưỡng cảnh báo và chính sách quyền riêng tư."),
        ("Bắt đầu phiên thi", "Thí sinh kiểm tra thông tin, chấp thuận thông báo và cấp quyền webcam."),
        ("Phân tích cục bộ", "MediaPipe xử lý khung hình; Behavior Engine tổng hợp tín hiệu theo thời gian."),
        ("Gửi sự kiện", "Máy thí sinh gửi trạng thái cùng cảnh báo cần chú ý, không truyền video liên tục."),
        ("Xác minh", "Giám thị kiểm tra trực tiếp, ghi nhận bối cảnh và xử lý theo quy chế của đơn vị."),
    ]
    for title, detail in workflow:
        add_numbered(document, WORKFLOW_ID, title, detail)
    add_callout(document, "Điểm kiểm soát", "Mọi sự kiện chỉ có giá trị như tín hiệu tham khảo cho đến khi được giám thị xác minh.", fill=PALE_GREEN, accent=GREEN)


def add_technology(document):
    document.add_page_break()
    document.add_heading("4. Công nghệ sử dụng", level=1)
    document.add_paragraph(
        "Dự án được xây dựng theo kiến trúc web hiện đại, ưu tiên xử lý AI tại thiết bị và triển khai đơn giản trên nền tảng đám mây."
    ).style = "Lead"

    table = document.add_table(rows=1, cols=3)
    headers = ["Lớp", "Công nghệ", "Mục đích sử dụng"]
    for index, text in enumerate(headers):
        p = table.rows[0].cells[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), size=10, color=NAVY, bold=True)
    tech_rows = [
        ("Giao diện", "Next.js 16, React 19, TypeScript", "Xây dựng ứng dụng web theo từng vai trò, kiểm soát kiểu dữ liệu và tạo bản production tối ưu."),
        ("Thiết kế", "Tailwind CSS 4, shadcn, Lucide React", "Tạo giao diện phản hồi tốt trên nhiều kích thước màn hình và hệ thống biểu tượng nhất quán."),
        ("AI thị giác", "MediaPipe Tasks Vision - Face Landmarker", "Phát hiện tối đa hai khuôn mặt và các điểm mốc phục vụ ước tính hướng mắt, quay, cúi và ngẩng đầu."),
        ("Thực thi AI", "WebAssembly, GPU delegate, Canvas API", "Chạy mô hình ngay trong trình duyệt, tăng tốc trên GPU và vẽ lớp điểm mốc lên video."),
        ("Webcam", "MediaDevices.getUserMedia", "Xin quyền và nhận luồng hình ảnh từ webcam qua môi trường HTTPS hoặc localhost."),
        ("Xử lý hành vi", "Behavior Engine bằng TypeScript", "Áp dụng ngưỡng thời gian, bộ đệm và khoảng nghỉ cảnh báo để hạn chế nhiễu."),
        ("Đồng bộ", "Firebase Firestore (tùy chọn)", "Đồng bộ trạng thái máy và sự kiện giữa nhiều thiết bị theo thời gian thực."),
        ("Demo cục bộ", "BroadcastChannel, localStorage", "Cho phép thử nghiệm nhiều tab trong cùng trình duyệt khi chưa cấu hình Firebase."),
        ("Triển khai", "Vercel, Node.js 24.x", "Dựng và phát hành ứng dụng Next.js qua HTTPS tại tên miền production."),
    ]
    for layer, technology, purpose in tech_rows:
        cells = table.add_row().cells
        for index, text in enumerate((layer, technology, purpose)):
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if index == 0:
                set_run_font(p.add_run(text), size=9.7, color=DARK_BLUE, bold=True)
            elif index == 1:
                set_run_font(p.add_run(text), size=9.7, color=NAVY, bold=True)
            else:
                set_run_font(p.add_run(text), size=9.7, color=BODY)
    configure_table(table, [1380, 2760, 5220])

    document.add_heading("4.1. Mô hình AI hiện tại", level=2)
    document.add_paragraph(
        "Mô hình chính là Google MediaPipe Face Landmarker, tệp face_landmarker.task dạng float16. "
        "Mô hình được tải trong lần chạy đầu tiên và thực hiện suy luận ở chế độ video ngay trên trình duyệt, ưu tiên GPU. "
        "Các chỉ số quay đầu, cúi/ngẩng và hướng mắt được tính từ điểm mốc khuôn mặt và đồng tử."
    )
    add_callout(
        document,
        "Lưu ý kỹ thuật",
        "Phiên bản hiện tại không sử dụng ChatGPT, mô hình sinh nội dung hoặc SCRFD trong luồng giám sát. SCRFD chỉ nên được xem xét bổ sung nếu thử nghiệm thực tế cho thấy cần tăng khả năng phát hiện khuôn mặt nhỏ, xa hoặc bị che khuất.",
        fill="FFF8E8",
        accent=GOLD,
    )

    document.add_heading("4.2. Kiến trúc dữ liệu", level=2)
    for title, detail in [
        ("Máy thí sinh", "Webcam → MediaPipe Face Landmarker → chỉ số khuôn mặt/hướng mắt/tư thế."),
        ("Bộ phân tích hành vi", "Chỉ số → ngưỡng thời gian → mức xanh/vàng/đỏ → sự kiện cảnh báo."),
        ("Kênh đồng bộ", "Sự kiện và trạng thái → Firebase Firestore; hoặc BroadcastChannel/localStorage trong chế độ demo."),
        ("Dashboard giám thị", "Nhận sự kiện → lọc và xem chi tiết → xác minh → xuất báo cáo."),
    ]:
        add_numbered(document, ARCHITECTURE_ID, title, detail)


def add_deployment_and_limits(document):
    document.add_heading("5. Triển khai và điều kiện vận hành", level=1)
    add_bullet(document, BULLET_ID, "Địa chỉ production: https://ai-examguard.vercel.app.")
    add_bullet(document, BULLET_ID, "Webcam chỉ hoạt động trên website HTTPS hoặc localhost và phải được người dùng cấp quyền.")
    add_bullet(document, BULLET_ID, "Nếu dùng nhiều máy khác nhau, cần cấu hình Firebase và các biến môi trường NEXT_PUBLIC_FIREBASE_* trên Vercel.")
    add_bullet(document, BULLET_ID, "Trước khi dùng thật phải bổ sung xác thực người dùng, phân quyền theo vai trò, Firebase App Check và Firestore Security Rules.")

    document.add_heading("5.1. Giới hạn cần lưu ý", level=2)
    add_bullet(document, BULLET_ID, "Độ chính xác phụ thuộc ánh sáng, vị trí webcam, chất lượng hình ảnh và đặc điểm người sử dụng.")
    add_bullet(document, BULLET_ID, "Các ngưỡng mặc định phục vụ trình diễn kỹ thuật; cần hiệu chỉnh bằng dữ liệu thử nghiệm của đơn vị tổ chức thi.")
    add_bullet(document, BULLET_ID, "Tín hiệu hướng mắt và tư thế đầu chỉ là ước tính từ hình ảnh 2D, không phải bằng chứng độc lập về hành vi vi phạm.")
    add_bullet(document, BULLET_ID, "Cần công khai mục đích giám sát, phạm vi dữ liệu, thời hạn lưu, quyền truy cập và quy trình khiếu nại trước khi triển khai.")
    add_callout(
        document,
        "Kết luận",
        "AI ExamGuard kết hợp AI thị giác chạy tại thiết bị, bộ quy tắc theo thời gian và dashboard theo vai trò để hỗ trợ giám sát minh bạch, có kiểm soát. Hiệu quả của hệ thống phụ thuộc vào quy trình kiểm thử, cấu hình phù hợp và sự xác minh của con người.",
        fill=PALE_GREEN,
        accent=GREEN,
    )


def build_document():
    global BULLET_ID, WORKFLOW_ID, ARCHITECTURE_ID
    document = Document()
    document.core_properties.title = "AI ExamGuard - Mô tả tính năng, vai trò và công nghệ sử dụng"
    document.core_properties.subject = "Tài liệu mô tả hệ thống AI hỗ trợ giám sát thi"
    document.core_properties.author = "AI ExamGuard"
    document.core_properties.keywords = "AI ExamGuard, MediaPipe, giám sát thi, Vercel, Next.js"
    configure_styles(document)
    configure_page(document)
    BULLET_ID = add_numbering_definition(document, kind="bullet", marker="•")
    WORKFLOW_ID = add_numbering_definition(document, kind="decimal", marker="%1.")
    ARCHITECTURE_ID = add_numbering_definition(document, kind="decimal", marker="%1.")

    add_cover(document)
    start_body_section(document)
    add_intro(document)
    add_features(document)
    add_roles(document)
    add_technology(document)
    add_deployment_and_limits(document)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
